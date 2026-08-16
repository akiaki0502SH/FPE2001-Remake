from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(r"C:\Users\猪\OneDrive\Documents\ChatGPT\FPE2001")
DELIVERABLES = ROOT / "deliverables"


def header_footer_paragraphs(document: Document):
    paragraphs = []
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in (*section.header.tables, *section.footer.tables):
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
    return paragraphs


def replace_runs(paragraphs, old: str, new: str) -> None:
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def replace_active_version(document: Document, old: str, new: str) -> None:
    # Only the cover title, active summary and page furniture are current-version
    # labels. Historical sections remain unchanged for implementation traceability.
    for index in (2, 5):
        if index < len(document.paragraphs):
            replace_runs([document.paragraphs[index]], old, new)
    if document.tables:
        cover_table_paragraphs = [
            paragraph
            for row in document.tables[0].rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
        replace_runs(cover_table_paragraphs, old, new)
    replace_runs(header_footer_paragraphs(document), old, new)


def append_section(document: Document, title: str, summary: str, sections: list[tuple[str, list[str]]]) -> None:
    document.add_page_break()
    document.add_paragraph(title, style="Heading 1")
    document.add_paragraph(summary)
    for heading, bullets in sections:
        document.add_paragraph(heading, style="Heading 2")
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")


def create_uiux() -> Path:
    source = DELIVERABLES / "FPE2001-Remake_UIUX实施规格_v2.6.docx"
    target = DELIVERABLES / "FPE2001-Remake_UIUX实施规格_v2.7.docx"
    copy2(source, target)
    document = Document(target)
    replace_active_version(document, "v2.6", "v2.7")
    append_section(
        document,
        "v2.7 Windows 11 原生 DWM 窗口框架",
        "本节替代 v2.6 的 WindowFrame Margin + DropShadowEffect 方案。窗口布局和操作逻辑保持不变；边框、四角裁切与窗口外阴影改由 Windows 11 Desktop Window Manager（DWM）绘制，视觉行为参考 Clash Verge 的原生窗口边界。",
        [
            (
                "问题更正",
                [
                    "不得在根内容外保留同色 Margin。根 Grid 直接铺满客户区，窗口外侧不能出现由 Window.Background 暴露形成的实色环带。",
                    "不得把 DropShadowEffect 挂到根 Border。WPF 效果位于窗口位图内部，会占用或污染客户区，不能视为真正的窗口外阴影。",
                    "不得用根 Border 的 ClipToBounds 模拟窗口圆角。四个外角必须由 DWM 对 HWND 表面进行系统级裁切。",
                ],
            ),
            (
                "视觉规格",
                [
                    "普通窗口使用 Windows 11 原生圆角（DWMWCP_ROUND）；四个角均应完整可见。最大化状态由系统取消圆角，与 Clash Verge 及 Windows 11 标准窗口一致。",
                    "活动窗口边框色为 #75869A，非活动窗口边框色为 #AAB6C4；边框由 DWMWA_BORDER_COLOR 绘制为系统/DPI 对齐的细线。",
                    "阴影完全位于窗口矩形之外，由 DWM 合成。阴影宽度与透明度跟随 Windows 11、DPI、激活状态及系统设置，不再用固定 2 DIP 的应用内效果模拟。",
                    "边框之外只允许出现可透出桌面内容的半透明阴影，不允许出现蓝灰色、白色或窗口背景色的实色边带。",
                ],
            ),
            (
                "交互与布局",
                [
                    "标题栏、九模块横向导航、动作栏、内容区与状态栏的尺寸和顺序不变；不为边框增加内容 Margin。",
                    "WindowChrome CaptionHeight=46、ResizeBorderThickness=6，GlassFrameThickness=1；最小化、最大化/还原和关闭按钮继续使用现有自定义按钮。",
                    "窗口拖拽、双击标题栏最大化、四边/四角缩放及按钮命中区域必须保持原行为。",
                ],
            ),
            (
                "验收标准",
                [
                    "普通窗口四角均可见且圆角边界连续；边框外无同色填充带，桌面背景可一直透到阴影下方。",
                    "窗口阴影位于 HWND 外部；移动窗口时阴影随系统重绘，客户区尺寸和控件位置不发生 2 DIP 偏移。",
                    "100%/125%/150%/200% DPI 下检查边界连续；最大化时边缘贴合工作区且不保留圆角或外边距。",
                    "最小化、最大化、还原、关闭、拖拽和调整大小全部通过回归测试。",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake UI/UX 实施规格 v2.7"
    document.core_properties.subject = "Windows 11 原生 DWM 边框、圆角与窗口外阴影规范"
    document.save(target)
    return target


def create_code_spec() -> Path:
    source = DELIVERABLES / "FPE2001-Remake_代码实施规格包_v2.3.docx"
    target = DELIVERABLES / "FPE2001-Remake_代码实施规格包_v2.4.docx"
    copy2(source, target)
    document = Document(target)
    replace_active_version(document, "v2.3", "v2.4")
    append_section(
        document,
        "v2.4 原生 DWM 窗口框架实施规格",
        "本节废止 v2.3 中根 Border(WindowFrame, Margin=2) 与 WindowFrameShadow 的实现要求，改用 Windows 11 非客户区合成。",
        [
            (
                "XAML 变更",
                [
                    "MainWindow 根元素恢复为原四行 Grid；删除窗口级 Border 的 Margin、BorderThickness、CornerRadius、ClipToBounds 与 Effect。",
                    "WindowChrome 使用 GlassFrameThickness=1、CaptionHeight=46、ResizeBorderThickness=6、UseAeroCaptionButtons=False；保留自定义窗口按钮与命中声明。",
                    "Colors.xaml 删除 WindowFrameShadow；WindowFrameBrush 仅作为活动 DWM 边框颜色资源，不再用于客户区 Border。",
                ],
            ),
            (
                "Win32/DWM 互操作",
                [
                    "新增 Shell/Windows11WindowFrame.cs，封装 DwmSetWindowAttribute；仅在 Windows 11 Build 22000 及以上执行。",
                    "SourceInitialized 后设置 DWMWA_NCRENDERING_POLICY=Enabled、DWMWA_WINDOW_CORNER_PREFERENCE=Round、DWMWA_BORDER_COLOR。",
                    "Activated 使用 WindowFrameBrush #75869A，Deactivated 使用 LineBrush #AAB6C4；COLORREF 按 0x00BBGGRR 转换。",
                    "所有 DWM 调用不进入 ViewModel，不改变命令、业务状态或页面生命周期；不支持的系统直接跳过属性设置。",
                ],
            ),
            (
                "禁止项",
                [
                    "禁止再次使用根级 DropShadowEffect、透明 Window、AllowsTransparency=True 或人为 2 DIP 空白模拟系统阴影。",
                    "禁止给根 Grid 增加为阴影预留的 Padding/Margin；窗口阴影不得参与 WPF Measure/Arrange。",
                    "禁止手工绘制四个外角遮罩；圆角状态由 DWM 根据普通/最大化状态管理。",
                ],
            ),
            (
                "测试契约",
                [
                    "Release 构建必须 0 警告、0 错误；既有单元测试全部通过。",
                    "带窗口外扩截图验证四角、边框与阴影：截图必须包含 GetWindowRect 外至少 12px 桌面区域。",
                    "UI Automation 必须分别验证最小化、最大化、还原、关闭；最大化后不得出现圆角或窗口外空白。",
                ],
            ),
            (
                "实现依据",
                [
                    "Microsoft Learn: DWMWINDOWATTRIBUTE - https://learn.microsoft.com/windows/win32/api/dwmapi/ne-dwmapi-dwmwindowattribute",
                    "Microsoft Learn: Apply rounded corners in desktop apps - https://learn.microsoft.com/windows/apps/desktop/modernize/ui/apply-rounded-corners",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake 代码实施规格包 v2.4"
    document.core_properties.subject = "Windows 11 原生 DWM 窗口框架实施契约"
    document.save(target)
    return target


def create_architecture() -> Path:
    source = DELIVERABLES / "FPE2001-Remake_Win11重构分析与架构设计_v2.3.docx"
    target = DELIVERABLES / "FPE2001-Remake_Win11重构分析与架构设计_v2.4.docx"
    copy2(source, target)
    document = Document(target)
    replace_active_version(document, "v2.3", "v2.4")
    append_section(
        document,
        "v2.4 Native DWM Frame 架构更正",
        "窗口边界从 WPF 客户区装饰提升为 Windows 11 非客户区系统能力。该更正消除实色外环、客户区缩进和圆角遮挡，并保持 UI Shell 以外的所有架构边界不变。",
        [
            (
                "职责边界",
                [
                    "WPF 客户区负责标题栏内容、模块导航、页面与状态栏；DWM 负责 HWND 外框、系统圆角及窗口外阴影。",
                    "Shell/Windows11WindowFrame 是唯一的原生边界适配器；领域层、扫描层、内存访问层、文件编辑层和 ViewModel 不引用 DWM。",
                    "WindowFrameBrush/LineBrush 提供颜色令牌，适配器将其转换为 COLORREF；视觉资源不泄漏为业务状态。",
                ],
            ),
            (
                "布局与合成原则",
                [
                    "客户区从 (0,0) 开始铺满窗口，不为系统阴影预留布局空间；外阴影不参与 WPF Measure/Arrange。",
                    "GlassFrameThickness=1 用于保留 DWM 非客户区合成，WindowChrome 继续提供拖拽、缩放和自定义标题栏命中。",
                    "普通窗口请求 DWMWCP_ROUND；最大化、贴靠和系统策略对圆角的最终决定权属于 Windows 11。",
                    "原生阴影的范围不是应用层固定像素契约；验收关注透明合成、无实色环带以及四角完整性。",
                ],
            ),
            (
                "兼容与失败模式",
                [
                    "Windows 11 Build 22000+ 启用边框颜色和圆角属性；低版本系统跳过这些属性并保留 WindowChrome 基础行为。",
                    "DwmSetWindowAttribute 返回失败时不得阻止主窗口创建；失败影响仅限系统装饰，不影响扫描、编辑或自动化功能。",
                    "激活状态变化只刷新 DWM 边框色，不重建窗口、不重载页面，也不触发业务命令。",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake Win11 重构分析与架构设计 v2.4"
    document.core_properties.subject = "Native DWM Frame 非客户区边界架构"
    document.save(target)
    return target


if __name__ == "__main__":
    outputs = (create_uiux(), create_code_spec(), create_architecture())
    for output in outputs:
        print(output)
