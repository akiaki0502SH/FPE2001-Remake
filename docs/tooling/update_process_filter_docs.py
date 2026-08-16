from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(r"C:\Users\猪\OneDrive\Documents\ChatGPT\FPE2001")
DELIVERABLES = ROOT / "deliverables"


def replace_runs(document: Document, old: str, new: str) -> None:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def append_section(document: Document, title: str, summary: str, sections: list[tuple[str, list[str]]]) -> None:
    document.add_page_break()
    document.add_paragraph(title, style="Heading 1")
    document.add_paragraph(summary)
    for heading, bullets in sections:
        document.add_paragraph(heading, style="Heading 2")
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")


def create_uiux() -> None:
    source = DELIVERABLES / "FPE2001-Remake_UIUX实施规格_v2.4.docx"
    target = DELIVERABLES / "FPE2001-Remake_UIUX实施规格_v2.5.docx"
    copy2(source, target)
    document = Document(target)
    replace_runs(document, "v2.4", "v2.5")
    append_section(
        document,
        "v2.5 进程列表优化",
        "本版本只优化扫描页默认目标进程列表的可见范围，不改变进程扫描、地址表、十六进制编辑器或模拟器适配器的操作逻辑。",
        [
            (
                "过滤边界",
                [
                    "隐藏 Session 0 中的 Windows 服务和非交互式系统进程。",
                    "隐藏已知 Windows 系统进程名，例如 System、svchost、services、lsass、csrss、winlogon、dwm 等。",
                    "隐藏可执行文件位于 %WINDIR% 目录树下的进程，避免系统组件充斥目标列表。",
                    "游戏、模拟器、模拟器后台进程和路径暂时无法读取的用户进程仍保留，避免误伤可调试目标。",
                ],
            ),
            (
                "交互反馈",
                [
                    "扫描页状态文本改为“可用目标进程”，并明确提示 Windows 系统进程已隐藏。",
                    "刷新操作继续使用原有“刷新”按钮；按 PID 打开进程内存视图的入口不受默认列表过滤影响。",
                    "排序规则保持进程名不区分大小写升序，同名进程按 PID 升序。",
                ],
            ),
            (
                "验收标准",
                [
                    "列表中不出现 Session 0 服务、已知系统进程或 %WINDIR% 下的可执行文件。",
                    "用户态游戏或模拟器仍能在列表中出现；进程路径不可读时不因路径缺失而被误隐藏。",
                    "进程在刷新期间退出不会导致列表刷新失败。",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake UI/UX 实施规格 v2.5"
    document.core_properties.subject = "进程列表过滤与深色边界蓝界面规范"
    document.save(target)


def create_code_spec() -> None:
    source = DELIVERABLES / "FPE2001-Remake_代码实施规格包_v2.1.docx"
    target = DELIVERABLES / "FPE2001-Remake_代码实施规格包_v2.2.docx"
    copy2(source, target)
    document = Document(target)
    replace_runs(document, "v2.1", "v2.2")
    append_section(
        document,
        "v2.2 进程目标过滤实施规格",
        "为扫描页增加可测试的 ProcessTargetFilter，默认列表只展示交互式用户目标；过滤发生在 Application 层，UI 继续消费原有进程 ID/名称契约。",
        [
            (
                "新增类型",
                [
                    "ProcessTargetDescriptor：ProcessId、ProcessName、SessionId、ExecutablePath。",
                    "ProcessTargetFilter.IsEligible：集中实现当前进程、Session 0、系统进程名和 Windows 目录路径判断。",
                ],
            ),
            (
                "枚举实现",
                [
                    "ScanApplicationService.ListTargetsAsync 使用 using 释放 Process 实例，读取路径失败时保留候选并依赖名称/会话规则判断。",
                    "进程刷新期间的 InvalidOperationException 和 Win32Exception 按单个进程跳过，不中断整个列表。",
                    "支持 CancellationToken，排序为进程名不区分大小写升序、PID 升序。",
                ],
            ),
            (
                "测试",
                [
                    "覆盖 Session 0 服务、已知系统进程、Windows 目录路径、游戏/模拟器、路径不可读和当前进程六类边界。",
                    "保持十六进制页按 PID 打开进程的能力，不把默认列表过滤误当作安全边界。",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake 代码实施规格包 v2.2"
    document.core_properties.subject = "目标进程过滤实施与测试契约"
    document.save(target)


def create_architecture() -> None:
    source = DELIVERABLES / "FPE2001-Remake_Win11重构分析与架构设计_v2.1.docx"
    target = DELIVERABLES / "FPE2001-Remake_Win11重构分析与架构设计_v2.2.docx"
    copy2(source, target)
    document = Document(target)
    replace_runs(document, "v2.1", "v2.2")
    append_section(
        document,
        "v2.2 目标进程可见性边界",
        "默认目标选择器采用“用户目标优先、系统进程隐藏”的展示策略，降低 Windows 11 环境下进程列表噪声，同时保留模拟器和无窗口后台目标的可调试性。",
        [
            (
                "架构决策",
                [
                    "过滤规则位于 Application 层，不下沉到 Win32 内存读写层，避免把 UI 展示策略与访问能力耦合。",
                    "过滤只影响默认枚举结果；用户通过十六进制编辑器输入 PID 时仍可显式打开目标。",
                    "路径检查使用目录边界匹配，避免把 C:\\Windows.old 等相邻目录误判为系统目录。",
                ],
            ),
            (
                "兼容性",
                [
                    "Session 0 作为服务隔离边界，适用于 Windows 11 常规桌面会话。",
                    "对受保护进程的路径访问失败采用保守保留策略，再由名称和会话规则进行过滤。",
                    "不依赖窗口句柄，因此无窗口模拟器后台进程仍可作为目标。",
                ],
            ),
        ],
    )
    document.core_properties.title = "FPE2001-Remake Win11 重构分析与架构设计 v2.2"
    document.core_properties.subject = "Windows 11 目标进程可见性边界"
    document.save(target)


if __name__ == "__main__":
    create_uiux()
    create_code_spec()
    create_architecture()
    print("created v2.5 UI/UX, v2.2 code specification, and v2.2 architecture documents")
