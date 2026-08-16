"""Bump the three specification documents to v2.9 / v2.6 / v2.6 and append the
v2.9 implementation section (auto scan type, candidate direct operations,
hex editor byte-level editing and locate-on-open).

Usage: python docs/tooling/update_v29_docs.py
"""

from pathlib import Path
from shutil import copy2

from docx import Document

SPEC = Path(__file__).resolve().parents[1] / "specifications"


def iter_all_paragraphs(document):
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
    return paragraphs


def replace_in_paragraph(paragraph, old, new):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)


def update_active_version(document, old, new):
    for index in (2, 5):
        if index < len(document.paragraphs):
            replace_in_paragraph(document.paragraphs[index], old, new)
    for paragraph in iter_all_paragraphs(document):
        replace_in_paragraph(paragraph, old, new)


def append_if_missing(document, title, summary, sections):
    if any(title in paragraph.text for paragraph in document.paragraphs):
        return
    document.add_page_break()
    document.add_paragraph(title, style="Heading 1")
    document.add_paragraph(summary)
    for heading, bullets in sections:
        document.add_paragraph(heading, style="Heading 2")
        for bullet in bullets:
            document.add_paragraph(bullet, style="List Bullet")


def bump(source_name, target_name, old, new, title, subject, section_title, summary, sections):
    source = SPEC / source_name
    target = SPEC / target_name
    if source.exists() and source.resolve() != target.resolve():
        copy2(source, target)
    document = Document(target)
    update_active_version(document, old, new)
    append_if_missing(document, section_title, summary, sections)
    document.core_properties.title = title
    document.core_properties.subject = subject
    document.save(target)
    print(f"created {target_name}")


UIUX_SECTIONS = [
    (
        "扫描（M01）",
        [
            "新增「自动」扫描类型并设为默认：首次扫描内存只读一遍，同时按 8/16/32 位整数三宽度判定，每宽度独立 FPSN 快照；再次扫描各快照独立过滤，自动收敛出真实类型与地址，解决模拟器游戏 8/16 位数值用 32 位搜不到的问题。",
            "候选列表新增「类型」列；首次/再次扫描后状态栏显示各宽度命中分布，0 候选时给出诊断提示（编码/类型、数值变化、进程实例选错）。",
            "值输入支持 0x 十六进制；值输入框与编辑器查找框按 Enter 直接触发。",
            "扫描热路径零分配优化：预编译期望值、span 化、unsafe fixed 零拷贝、写缓冲与块读、宽度特化热循环、进度节流；64 MiB 全内存自动扫描由 11.9 s 降至 0.23 s。",
        ],
    ),
    (
        "扫描结果直接操作（原版 FPE2001 能力）",
        [
            "「写入」：输入框改值 → 自动建立地址表条目 → 立即写内存（写前比较 + 读回校验），候选行当前值同步。",
            "「锁定」：持续写回选中候选值（Always，250 ms，失败熔断），再点解除；条目同步进入地址表。",
            "「在编辑器中打开」：跳转十六进制编辑器并定位到该地址。",
        ],
    ),
    (
        "十六进制编辑器（M03）",
        [
            "字节级选中：十六进制区拆为 16 个独立字节列（00–0F）+ SelectionUnit=Cell，点击单个字节黄色高亮（#FFD54F），编辑框预填该字节 hex。",
            "打开即定位：绝对地址 → 视图内偏移映射；目标行滚到可视区第一行、目标字节列滚到最左列（左上角第一位）并高亮。",
            "进程内存「应用覆盖」即写回目标进程；写前比较失败自动刷新视图提示重试。",
            "修复 Cell 选择模式下 ScrollIntoView 选择整行导致的崩溃；新增全局异常兜底与 crash 日志。",
        ],
    ),
]

CODE_SECTIONS = [
    (
        "扫描引擎（ScanEngine）",
        [
            "SessionState 扩展为多快照：自动模式同时维护 UInt8/UInt16/UInt32 三个 FPSN 快照，再次扫描逐快照过滤；ScanSession 附带各快照类型与候选数统计。",
            "ValueComparer 预编译期望值（ExpectedValue 结构）并在循环外编译一次；Matches/MatchesFirst 全部 ReadOnlySpan 化，消除每候选数组分配。",
            "ProcessHandle.Read/Write 改为 unsafe fixed 零拷贝 P/Invoke（byte*），消除每块 ToArray 拷贝。",
            "FpsnSnapshot 增加 256 KiB 写缓冲与 ReadAllBulk 块读迭代（回调可提前终止）；进度报告 200 ms 节流。",
            "自动模式 ScanRegionAuto：单遍读内存 + 宽度特化热循环（8/16/32 位全偏移判定，跨块 tail 重叠处理）。",
            "ReadBlockWithRetry：再次扫描块读失败逐级缩小重试，避免小区域/区域边界候选被整块丢弃。",
        ],
    ),
    (
        "十六进制编辑器（HexEditorViewModel / VirtualHexRows）",
        [
            "VirtualHexRows 增加行实例缓存：同一索引返回同一实例，保证 DataGrid.ScrollIntoView/选中引用一致（修复定位失败）。",
            "HexRow 增加字节索引器 this[int] 与 16 列绑定；SelectedByteOffset 字节级编辑起点。",
            "OpenProcessAtCoreAsync 做绝对地址 → 视图内偏移映射并越界防御；ScrollToOffset 记录 FocusByteColumn 与 pending 定位（页面就绪后消费）。",
            "进程内存来源「应用覆盖」自动提交（VerifyVersionToken=false 路径）并重载视图；写前比较失败自动刷新。",
        ],
    ),
    (
        "UI 与异常处理",
        [
            "ScanViewModel 注入 IFreezeService 与编辑器跳转回调；候选写入/锁定复用地址表服务，创建条目后立即写入/冻结。",
            "全局 DispatcherUnhandledException 兜底：记录 crash 日志到 %TEMP%\\FPE2001-Remake\\logs 并提示，不再无提示闪退。",
        ],
    ),
]

ARCH_SECTIONS = [
    (
        "扫描会话的多类型扩展",
        [
            "自动扫描把会话从单一 DataType/快照扩展为多快照（8/16/32 位各一），共享同一进程句柄与区域规划；取消语义与原子替换保持不变。",
            "候选展示按 32→16→8 顺序合并；快照统计随 ScanSession 返回，供 UI 呈现各宽度命中分布。",
        ],
    ),
    (
        "地址模型与编辑器映射",
        [
            "扫描候选持有 HostVirtual 绝对地址；编辑器打开 4KB 页对齐窗口后必须转换为视图内偏移，绝对地址直接定位会越界（已修复并加防御）。",
            "定位链路：绝对地址 → 文档偏移 → 行/字节列 → ScrollIntoView + ScrollTo 精确对齐（行到顶、列到最左）+ 单元格高亮。",
            "进程内存写入走 IEditableByteSource.CommitAsync 等长覆盖，写前比较 + 读回校验保证不覆盖竞态写入。",
        ],
    ),
]


def create_uiux():
    bump(
        "FPE2001-Remake_UIUX实施规格_v2.9.docx",
        "FPE2001-Remake_UIUX实施规格_v2.9.docx",
        "v2.8",
        "v2.9",
        "FPE2001-Remake UI/UX 实施规格 v2.9",
        "扫描自动类型判断、扫描结果直接操作与十六进制字节级编辑定位规范",
        "v2.9 扫描自动类型与结果直接操作",
        "本版本为扫描工作区补齐原版 FPE2001 的搜索结果直接操作能力（写入/锁定/打开内存），并新增自动类型扫描与十六进制编辑器字节级定位编辑。",
        UIUX_SECTIONS,
    )


def create_code_spec():
    bump(
        "FPE2001-Remake_代码实施规格包_v2.6.docx",
        "FPE2001-Remake_代码实施规格包_v2.6.docx",
        "v2.5",
        "v2.6",
        "FPE2001-Remake 代码实施规格包 v2.6",
        "扫描引擎自动模式、零分配热路径与编辑器字节级编辑的实施契约",
        "v2.6 扫描引擎自动模式与编辑器字节级编辑",
        "本版本把扫描会话扩展为多快照自动模式，全面优化扫描热路径分配，并实现十六进制编辑器字节级选中与打开即定位。",
        CODE_SECTIONS,
    )


def create_architecture():
    bump(
        "FPE2001-Remake_Win11重构分析与架构设计_v2.6.docx",
        "FPE2001-Remake_Win11重构分析与架构设计_v2.6.docx",
        "v2.5",
        "v2.6",
        "FPE2001-Remake Win11 重构分析与架构设计 v2.6",
        "自动扫描多快照会话与地址映射/编辑器定位的架构边界",
        "v2.6 自动扫描会话扩展与地址映射边界",
        "本版本补充自动扫描多快照会话的架构说明，以及扫描候选绝对地址与编辑器视图内偏移的映射规则。",
        ARCH_SECTIONS,
    )


if __name__ == "__main__":
    create_uiux()
    create_code_spec()
    create_architecture()
    print("created UI/UX v2.9, code specification v2.6, and architecture v2.6 documents")
