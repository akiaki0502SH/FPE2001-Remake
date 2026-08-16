#!/usr/bin/env python3
from __future__ import annotations

import datetime as dt
import hashlib
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
WORK_DIR = ROOT / "artifacts_work"
OUT = OUT_DIR / "FPE2001_Win11重构分析与架构设计.docx"
ARCH_IMG = WORK_DIR / "fpe2001_target_architecture.png"

FONT = "Microsoft YaHei"
FONT_UI = "Microsoft YaHei UI"
MONO = "Consolas"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "2C6E49"
BORDER = "B7C3D0"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    tags = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for tag in tags:
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_num_definitions(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def create_abstract(kind, abs_id):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(3):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            lvl.append(start)
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
            lvl.append(num_fmt)
            lvl_text = OxmlElement("w:lvlText")
            marker = ["•", "–", "◦"][level] if kind == "bullet" else f"%{level + 1}."
            lvl_text.set(qn("w:val"), marker)
            lvl.append(lvl_text)
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            lvl.append(jc)
            ppr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            left = 540 + level * 540
            tab.set(qn("w:pos"), str(left))
            tabs.append(tab)
            ppr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(left))
            ind.set(qn("w:hanging"), "270")
            ppr.append(ind)
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:after"), "80")
            spacing.set(qn("w:line"), "300")
            spacing.set(qn("w:lineRule"), "auto")
            ppr.append(spacing)
            lvl.append(ppr)
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), FONT)
            rfonts.set(qn("w:hAnsi"), FONT)
            rfonts.set(qn("w:eastAsia"), FONT)
            rpr.append(rfonts)
            lvl.append(rpr)
            abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(numbering.index(first_num), abstract)

    def create_num(abs_id, num_id, restart=False):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abs_id))
        num.append(abs_ref)
        if restart:
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), "0")
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), "1")
            override.append(start_override)
            num.append(override)
        numbering.append(num)

    create_abstract("bullet", next_abs)
    create_abstract("decimal", next_abs + 1)
    create_num(next_abs, next_num)
    decimal_ids = []
    for offset in range(1, 4):
        create_num(next_abs + 1, next_num + offset, restart=True)
        decimal_ids.append(next_num + offset)
    return next_num, decimal_ids


def apply_num(paragraph, num_id, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(numid)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT)
    rpr.append(rfonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, text, end])
    set_run_font(run, FONT, 9, color=MUTED)


def add_para(doc, text="", style=None, size=None, bold=False, italic=False, color=None, align=None, before=0, after=None, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, FONT, size, bold=bold, italic=italic, color=color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text, bullet_num_id, level=0, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id, level)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, FONT, 11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, FONT, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, FONT, 11)
    return p


def add_number(doc, text, number_num_id, level=0):
    p = doc.add_paragraph()
    apply_num(p, number_num_id, level)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, FONT, 11)
    return p


def add_callout(doc, label, text, fill=CALLOUT, label_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120)
    set_table_borders(table, color=BORDER, size=6, inside=False)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(label + "  ")
    set_run_font(r1, FONT, 10.5, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, FONT, 10.5, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths, font_size=9.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, indent=120)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(text))
        set_run_font(r, FONT, font_size, bold=True, color=INK)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "FAFBFC")
            p = cells[i].paragraphs[0]
            p.alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(str(value))
            set_run_font(r, FONT, font_size, color="20262D")
        set_table_geometry(table, widths, indent=120)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    for line_no, line in enumerate(text.splitlines()):
        if line_no:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, MONO, 9, color="24303A")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, FONT, {1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return p


def draw_architecture():
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1400, 840), "#FFFFFF")
    d = ImageDraw.Draw(img)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    font_bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    base = ImageFont.truetype(str(font_path), 28) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 22) if font_path.exists() else ImageFont.load_default()
    bold = ImageFont.truetype(str(font_bold_path if font_bold_path.exists() else font_path), 30) if font_path.exists() else ImageFont.load_default()

    def box(x1, y1, x2, y2, title, lines, fill, outline=BLUE):
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#" + fill.lstrip("#"), outline="#" + outline.lstrip("#"), width=3)
        d.text((x1 + 18, y1 + 14), title, font=bold, fill="#" + INK)
        y = y1 + 58
        for line in lines:
            d.text((x1 + 20, y), line, font=small, fill="#27323B")
            y += 34

    def arrow(x1, y1, x2, y2):
        d.line((x1, y1, x2, y2), fill="#607487", width=4)
        if y2 >= y1:
            pts = [(x2, y2), (x2 - 9, y2 - 18), (x2 + 9, y2 - 18)]
        else:
            pts = [(x2, y2), (x2 - 9, y2 + 18), (x2 + 9, y2 + 18)]
        d.polygon(pts, fill="#607487")

    d.text((40, 26), "FPE Next：Win11 x64 目标架构", font=bold, fill="#" + INK)
    box(70, 100, 1330, 210, "桌面前端（WPF / .NET 10 x64）", ["进程与模拟器选择  ·  扫描工作流  ·  地址表  ·  内存查看器  ·  诊断中心"], "F7FAFD")
    arrow(700, 212, 700, 252)
    box(70, 260, 1330, 390, "应用服务层", ["会话编排  ·  权限提示  ·  任务取消  ·  配置与数据库  ·  审计日志  ·  旧格式导入"], "EEF4FA")
    arrow(700, 392, 700, 430)
    box(70, 438, 805, 610, "64 位核心引擎", ["区域枚举与读取", "初扫/再扫/结果快照", "类型与字节序", "条件写入/冻结/回滚"], "E8EEF5")
    box(835, 438, 1330, 610, "适配器主机（隔离进程）", ["能力协商", "Guest ↔ Host 地址翻译", "版本探测与签名定位", "诊断包与可回放测试"], "FFF8E8", outline=GOLD)
    arrow(700, 612, 700, 654)
    box(70, 665, 435, 805, "通用 Win32 后端", ["VirtualQueryEx", "Read/WriteProcessMemory", "x86 / x64 目标"], "F4F6F9")
    box(520, 665, 885, 805, "模拟器原生接口", ["调试 API / IPC", "内存域 / 暂停 / 帧同步", "优先于硬编码偏移"], "F4F6F9")
    box(970, 665, 1330, 805, "Libretro / 专用桥接", ["System RAM / Save RAM", "核心版本与域描述", "逐核心验证"], "F4F6F9")
    for x in (250, 700, 1150):
        arrow(x, 655, x, 667)
    img.save(ARCH_IMG, dpi=(180, 180))


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb("20262D")
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in (
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    code = doc.styles.add_style("Code Block", 1)
    code.font.name = MONO
    code.font.size = Pt(9)
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05
    ppr = code._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    ppr.append(shd)


def configure_header_footer(doc):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = ""
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("FPE 2001 / Win11 Reconstruction")
        set_run_font(r, FONT, 8.5, color=MUTED)

        footer = section.footer
        p0 = footer.paragraphs[0]
        p0._element.getparent().remove(p0._element)
        t = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(t, [7000, 2360], indent=0)
        remove_table_borders(t)
        lp = t.cell(0, 0).paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lr = lp.add_run("FPE 2001 Win11 重构设计 · 静态分析版")
        set_run_font(lr, FONT, 8.5, color=MUTED)
        rp = t.cell(0, 1).paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rr = rp.add_run("第 ")
        set_run_font(rr, FONT, 8.5, color=MUTED)
        add_page_field(rp)
        rr2 = rp.add_run(" 页")
        set_run_font(rr2, FONT, 8.5, color=MUTED)


def add_cover(doc):
    add_para(doc, "TECHNICAL RECONSTRUCTION REPORT", size=10, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=62, after=24)
    add_para(doc, "FPE 2001", size=34, bold=True, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "Win11 重构分析与架构设计", size=23, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_para(doc, "面向 64 位进程、超 4GB 地址与多模拟器逐个适配", size=13.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=70)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2300, 7060], indent=120)
    set_table_borders(table, color="D8DEE5", size=5)
    rows = [
        ("分析对象", "FPE2001_绿色版.zip（静态分析，不执行样本）"),
        ("目标平台", "Windows 11 x64；第一阶段不承诺 ARM64 原生"),
        ("目标用途", "本地、离线、单机游戏与模拟器调试；不绕过反作弊或受保护进程"),
        ("文档版本", "v0.1 · 2026-08-15"),
    ]
    for i, (label, value) in enumerate(rows):
        set_cell_shading(table.cell(i, 0), LIGHT_BLUE)
        p1 = table.cell(i, 0).paragraphs[0]
        r1 = p1.add_run(label)
        set_run_font(r1, FONT, 10, bold=True, color=INK)
        p2 = table.cell(i, 1).paragraphs[0]
        r2 = p2.add_run(value)
        set_run_font(r2, FONT, 10, color="20262D")
    add_para(doc, "结论先行：建议重写，不建议在原 32 位二进制上继续打补丁。", size=11.5, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=42, after=10)
    add_para(doc, "核心思路：把“模拟器差异”从扫描器中剥离，变成可诊断、可测试、可独立升级的适配器。", size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def add_contents(doc, number_num_id):
    add_heading(doc, "阅读导航", 1)
    items = [
        "结论与优先级",
        "样本与静态分析边界",
        "旧版前端功能还原",
        "旧版后端与数据流",
        "Win11 重构目标与非目标",
        "目标架构与技术选型",
        "64 位地址模型",
        "扫描、写入与冻结引擎",
        "模拟器适配器协议",
        "逐模拟器调试流程",
        "前端信息架构与旧功能迁移",
        "安全、测试、路线图与验收",
        "附录：哈希、配置与参考资料",
    ]
    for item in items:
        add_number(doc, item, number_num_id)
    add_callout(doc, "文档定位", "这是可直接进入方案评审和任务拆分的重构基线，不是对每个模拟器已经完成兼容的承诺。每个模拟器仍需按版本逐个校准与回归。")


def add_body(doc, bullet_num_id, number_num_ids):
    strategy_num_id, steps_num_id = number_num_ids[1], number_num_ids[2]
    add_heading(doc, "1. 结论与优先级", 1)
    add_callout(doc, "推荐决策", "保留产品意图和用户工作流，放弃原二进制、旧运行库、全局钩子实现和硬编码地址。新建 x64 核心、最低权限内存代理和独立模拟器适配器主机。")
    for text in [
        "原版已经具备完整工具雏形：进程选择、初扫/再扫、8/16/32 位数值、结果筛选、地址表、立即写入、冻结/条件写入、十六进制编辑、搜索任务快照、键鼠宏、变速、图片查看与导出。",
        "不能直接“扩成 64 位”：FPE.exe、SPE.exe、tt.dll 均为 PE32；地址格式大量使用 %08X，界面上限为 4GB，核心通过 32 位指针调用 VirtualQueryEx、ReadProcessMemory、WriteProcessMemory。",
        "模拟器支持必须以地址空间抽象为中心：用户看到的是 Guest 地址，Windows API 操作的是 Host 虚拟地址；二者可能存在基址、镜像、分段、字节序和版本差异。",
        "MVP 先完成通用 Win32/x64 后端与测试靶进程，再做“旧 INI 兼容适配器”，随后由用户选择一个现代模拟器作为首个正式适配器。",
        "写入功能默认关闭；禁止内核驱动、代码注入和反作弊绕过。受保护进程或权限不匹配时应明确拒绝，而不是尝试规避。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "1.1 建议技术栈", 2)
    add_table(doc,
              ["层", "建议", "原因"],
              [
                  ("桌面 UI", ".NET 10 LTS + WPF，x64，自包含发布", "成熟的桌面控件、虚拟化列表、高 DPI 与传统 Win32 互操作；部署比 WinUI 3 更直接。"),
                  ("核心引擎", "C# 14 / .NET 10，Span<T>、Channels、SIMD", "足以覆盖扫描吞吐、异步取消和 64 位地址；先避免多语言 FFI。"),
                  ("Windows 互操作", "LibraryImport + SafeHandle", "显式匹配 Win32 类型和句柄生命周期，统一 GetLastError 诊断。"),
                  ("持久化", "SQLite + JSON 适配器清单", "结构化会话、迁移版本、可查询审计；适配器描述可审阅。"),
                  ("进程隔离", "命名管道 IPC；适配器独立进程", "UI 不因适配器崩溃而退出；权限边界和协议版本清晰。"),
                  ("安装", "签名的 x64 自包含包；后续可加 MSIX", "第一阶段降低部署摩擦，保留代码签名与升级通道。"),
              ], [1700, 2600, 5060])

    add_heading(doc, "2. 样本与静态分析边界", 1)
    add_para(doc, "本次仅对用户提供的压缩包做解包、哈希、PE 头、导入/导出表、资源窗体（Binary DFM）、可见字符串和配置文件分析；没有运行 FPE.exe、SPE.exe 或加载 tt.dll。")
    add_table(doc,
              ["证据级别", "含义", "本文示例"],
              [
                  ("已确认", "可由文件结构、导入表、窗体资源或配置直接复核", "32 位 PE；Read/WriteProcessMemory；地址表列；Fpe.ini 六条模拟器记录。"),
                  ("高置信推断", "多项静态证据一致，但未做运行时观察", "FPEMEM*.DAT 用作扫描快照；tt.dll 用于键鼠钩子与宏；锁定线程按周期写值。"),
                  ("待验证", "需要样例文件、运行日志或源代码", ".fpe/.mis 的精确二进制格式；INI 两个十六进制字段的全部语义；各模拟器映射稳定性。"),
              ], [1700, 3550, 4110])
    add_callout(doc, "重要限制", "“前端/后端”是按现代架构重新归类。原程序是单体桌面应用，并不存在 Web 前端、HTTP 服务或独立服务器。", fill="FFF8E8", label_color=GOLD)

    add_heading(doc, "2.1 文件角色", 2)
    add_table(doc,
              ["文件", "静态结论", "重构处理"],
              [
                  ("FPE.exe", "1,449,984 字节；PE32 x86；VCL/C++Builder 风格；主界面与内存工具。", "仅作为行为参考，不复用二进制或代码。"),
                  ("SPE.exe", "PE32 x86；标题为 Show Picture Expert；浏览图片并导出 BMP/JPG/GIF，含缩放窗口。", "作为可选“图片工具”后置，不进入内存 MVP。"),
                  ("tt.dll", "PE32 x86；导出 Setup/UnSetup/Setkey；使用 SetWindowsHookEx、keybd_event、mouse_event，共享节约 128KB。", "不复用。宏系统改为显式、可撤销、无全局注入的实现。"),
                  ("Mfc42.dll / Msvcrt.dll", "VC6 时代运行库；无签名；随包私带。", "删除依赖，使用受支持运行时和签名发布。"),
                  ("Fpe.ini", "包含六个旧模拟器记录，格式与程序字符串 %d,%X,%X,%s 对应。", "提供只读导入器，转换成版本化适配器配置。"),
              ], [1800, 4480, 3080])

    add_heading(doc, "3. 旧版前端功能还原", 1)
    add_para(doc, "主窗体尺寸为 632×449，使用 9 个标签页（由图标显示，资源名可还原其职责）。下表按资源控件、菜单项和事件名归纳。")
    add_table(doc,
              ["区域", "已确认的交互", "迁移建议"],
              [
                  ("扫描（TabScan）", "目标进程下拉、扫描值、8/16/32 位、Scan/Stop、Mission 1、F2–F11 快捷操作、候选地址清单。", "拆成连接栏、扫描条件、结果虚拟列表和会话历史。"),
                  ("地址表（TabTabs）", "列为 Address / Lock Value / if / Comment；添加、删除、修改、立即写入、加载/保存 .fpe。", "保留；地址表达式升级为 64 位、Guest/Host 域、条件规则和写入策略。"),
                  ("编辑器（TabEditor）", "18 列内存/十六进制网格，Hex 与 10/str 模式，打开、查找、撤销、刷新。", "实现虚拟化内存查看器；默认只读，编辑时显式提交。"),
                  ("文件选择/扫描（TabSelect）", "目录、驱动器、过滤器；列表列 Last time / File name / Size；扫描、编辑保存、刷新、停止。", "与内存核心解耦；如保留，作为单独文件工具插件。"),
                  ("GPE 图片区（TabGPE）", "ON/OFF、路径、文件名、序号、x2、预览以及加载/编辑/重命名/删除。", "归入可选媒体工具，不阻塞第一版。"),
                  ("键鼠宏（TabMacKey）", "9 行宏表；设置、删除、更新；加入左键、右键、双击；存在 Ctrl+F8/F11/F12 自动鼠标动作。", "改为前台应用范围热键；显示录制状态、倒计时和紧急停止键。"),
                  ("变速（TabSpeed）", "开关、0–100 滑杆、地址列表与重定位事件。", "不做通用进程“加速”；只在适配器声明 frame/rate 能力时开放。"),
                  ("设置（TabOthers）", "内存上限 Auto/2GB/3GB/4GB/DOS-16bit、临时目录、扫描/数据库/SPE/Undo 开关、锁定周期。", "改为设置中心；地址上限由目标进程自动探测，不再由用户选 4GB。"),
                  ("关于", "作者链接和邮件。", "替换为版本、许可证、诊断与隐私说明。"),
              ], [1900, 4700, 2760], font_size=9.1)

    add_heading(doc, "3.1 独立图片程序 SPE", 2)
    add_para(doc, "SPE 的窗体标题为“Show Picture Expert - by Guoo-jaw Li”，具备文件列表、图像预览、缩放窗口、调用画图程序，以及按质量导出 JPG、BMP、GIF。它不导入进程内存 API，因此应视为附属图片工具，而不是扫描后端。")

    add_heading(doc, "4. 旧版后端与数据流", 1)
    add_table(doc,
              ["子系统", "静态证据", "可能的数据流"],
              [
                  ("进程附加", "OpenProcess、EnumWindows、GetWindowThreadProcessId", "窗口/进程选择 → PID → 进程句柄。"),
                  ("区域与扫描", "VirtualQueryEx、ReadProcessMemory、Scan:(%x-%x)、FPEMEM01..DAT", "枚举可读区域 → 分块读取 → 写候选地址/旧值快照 → 再扫过滤。"),
                  ("写入与锁定", "WriteProcessMemory、TLockThread、LockTrackBar、条件列 if", "解析地址 → 单次写入或后台周期写入 → 条件判断/宏。"),
                  ("地址与类型", "%08X、8/16/32、float 格式、Can't calculate address", "32 位地址解析、整数/浮点显示和简单地址计算。"),
                  ("会话与表", ".fpe、.mis、.mi0、Mission、Undo", "地址表与扫描任务保存/恢复，撤销文件驻留临时目录。"),
                  ("偏好设置", "SOFTWARE\\jaw\\fpe、INI、临时路径、热键", "注册表与 INI 混合保存。"),
                  ("键鼠支持", "tt.dll：Setup/UnSetup/Setkey、Windows hooks、keybd_event/mouse_event", "全局钩子捕获热键，线程/共享节驱动宏。"),
              ], [1700, 3900, 3760], font_size=9.2)

    add_heading(doc, "4.1 Fpe.ini 的旧模拟器映射", 2)
    add_code(doc, "[Emulator]\nCount=6\nE01=1,5AE860,0,ePSXe_1.20\nE02=1,2C20180,0,FPSE_0.09\nE03=1,4010000,0,Connectix_VGS_1.41\nE04=1,0,A00000,bleem_1.5b\nE05=0,CA4240,0,Snes9x_1.36\nE06=0,FFC2DF00,0,ZSNESWIN_1.29r1.0")
    add_para(doc, "记录格式与程序内字符串“%d,%X,%X,%s”一致。第 1 字段很可能是启用/模式，第 2、3 字段是地址转换参数，第 4 字段是显示名；但精确变换公式必须通过运行时样例验证，文档不把推断写成事实。")

    add_heading(doc, "5. 为什么必须重写", 1)
    issues = [
        ("地址天花板", "PE32、%08X 和 4GB 选项贯穿界面与数据模型，无法可靠表示 x64 模拟器高地址。"),
        ("安全基线", "三个自有 PE 均未启用 ASLR/NX；所有可执行文件未签名；附带 VC6 运行库。"),
        ("模块耦合", "扫描、UI、文件工具、图片、宏、热键和配置集中在同一进程，任一故障可能拖垮全局。"),
        ("硬编码映射", "模拟器以版本名和固定十六进制参数存在 INI 中，升级后容易静默写错地址。"),
        ("权限过宽", "旧式工具通常一次申请较大权限；Win11 上更容易遇到完整性级别、DACL、受保护进程和安全软件拦截。"),
        ("不可测试", "无地址域、适配器能力和版本化格式，难以对“同一游戏、不同模拟器版本”做可重复回归。"),
        ("数据脆弱", "扫描快照依赖临时 DAT 文件，进程重启、地址空间变化或异常退出都可能产生陈旧结果。"),
    ]
    add_table(doc, ["问题", "影响"], issues, [1900, 7460])

    add_heading(doc, "6. Win11 重构目标与非目标", 1)
    add_heading(doc, "6.1 必须实现", 2)
    for text in [
        "以 x64 进程运行，完整保存和显示 64 位 Host 地址；持久化使用 UInt64，不把地址存成 Int32、字符串截断或当前进程 IntPtr。",
        "区分 HostVirtual、GuestPhysical、GuestVirtual、ModuleRelative 和 PointerChain 地址空间。",
        "支持 x86 与 x64 目标；附加后用 IsWow64Process2 记录目标架构和原生架构。",
        "扫描任务可取消、可暂停、可恢复；大型结果集磁盘化，UI 列表虚拟化。",
        "模拟器适配器具备版本探测、内存域、地址翻译、能力协商、健康检查和诊断导出。",
        "每个模拟器独立发布与回归；核心版本更新不要求同时修改全部适配器。",
        "默认只读；任何写入、冻结、宏都必须由用户明确开启，并可一键停止。",
    ]:
        add_bullet(doc, text, bullet_num_id)
    add_heading(doc, "6.2 明确不做", 2)
    for text in [
        "不绕过反作弊、DRM、受保护进程、内核保护或第三方安全边界。",
        "不提供内核驱动、远程线程注入、DLL 注入或隐藏进程行为。",
        "不保证一个硬编码偏移跨模拟器版本长期有效。",
        "不在 MVP 重做 SPE/GPE、通用变速和全局键鼠录制；这些属于后续插件。",
        "不直接写出未验证的旧 .fpe/.mis 文件，避免破坏用户数据。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "7. 目标架构", 1)
    draw_architecture()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(ARCH_IMG), width=Inches(6.45))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", "FPE Next 的分层架构：WPF 前端、应用服务、64 位核心、隔离适配器主机及三类后端。")
    p.paragraph_format.space_after = Pt(4)
    cap = add_para(doc, "图 1  Win11 x64 目标架构", size=9, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    cap.paragraph_format.keep_with_next = False

    add_heading(doc, "7.1 进程边界", 2)
    add_table(doc,
              ["进程", "默认权限", "职责", "失败影响"],
              [
                  ("FpeNext.UI.exe", "普通用户", "界面、任务编排、结果展示、配置", "仅 UI 会话中断；不持有长期写句柄。"),
                  ("FpeNext.Broker.exe", "按需、尽量不提权", "Win32 进程句柄、区域枚举、批量读写", "关闭目标句柄；任务可重新附加。"),
                  ("FpeNext.AdapterHost.exe", "普通用户、每适配器实例", "加载适配器、探测版本、翻译地址", "单适配器重启，不影响 UI 与其他适配器。"),
              ], [1900, 1700, 3300, 2460], font_size=9.2)
    add_para(doc, "Broker 首先以最小权限打开：只读扫描请求 PROCESS_QUERY_INFORMATION | PROCESS_VM_READ；只有写入会话才追加 PROCESS_VM_WRITE | PROCESS_VM_OPERATION。不得默认申请 PROCESS_ALL_ACCESS。")

    add_heading(doc, "8. 64 位地址模型", 1)
    add_code(doc, "public enum AddressSpace {\n    HostVirtual, GuestPhysical, GuestVirtual, ModuleRelative, PointerChain\n}\n\npublic readonly record struct LogicalAddress(\n    AddressSpace Space,\n    string DomainId,        // 例如 psx.ram / snes.wram / host.private\n    ulong Value,            // 永不使用 uint 保存持久地址\n    Endianness Endian,\n    int PointerWidthBits);\n\npublic readonly record struct ResolvedAddress(\n    ulong HostVirtualAddress,\n    int ProcessId,\n    long ProcessStartTime,  // 防止 PID 复用\n    string AdapterId,\n    string ResolutionProof  // 基址/签名/版本等诊断信息\n);")
    add_heading(doc, "8.1 关键规则", 2)
    for text in [
        "UI 默认显示 0x + 16 位十六进制；允许隐藏前导零，但复制、日志、导出保留完整值。",
        "Guest 地址与 Host 地址永不混用一个裸 ulong；所有 API 同时携带 AddressSpace 与 DomainId。",
        "指针链每一级声明读取宽度与字节序；x86 目标通常 32 位指针，x64 目标 64 位，不能由工具自身位数推断。",
        "模块相对地址保存 module identity + RVA，不保存某次启动的模块基址。",
        "适配器翻译结果绑定 PID + 启动时间 + 内存映射代次；进程重启或载入新 ROM 后自动失效。",
        "旧 INI 中的两个十六进制参数暂存为 LegacyTransformArg1/2，只有适配器验证后才转换成正式映射规则。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "9. 扫描、写入与冻结引擎", 1)
    add_heading(doc, "9.1 区域枚举与读取", 2)
    for text in [
        "用 VirtualQueryEx 从 0 递增到目标最大用户地址；只扫描 MEM_COMMIT 且可读的页面，跳过 PAGE_NOACCESS 与 PAGE_GUARD。",
        "按区域边界和固定块大小（建议 1–8 MiB 可调）切分读取。ReadProcessMemory 要求请求范围可访问，因此不可跨不可访问页面盲读。",
        "读取失败必须记录 Win32 错误码、区域属性、地址和长度；允许局部任务继续，但 UI 明确显示“跳过区域”，不能静默丢失。",
        "区域在扫描期间可能释放或变更；每个块使用版本戳和任务取消令牌，结果提交前校验目标进程仍是同一实例。",
        "缓冲区使用 ArrayPool<byte>；CPU 比较流水线与进程读取解耦，避免并发过高反而放大上下文切换。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "9.2 数据类型与筛选状态", 2)
    add_table(doc,
              ["类别", "MVP", "后续"],
              [
                  ("整数", "Int8/16/32/64、UInt8/16/32/64", "BCD、定点数、自定义位域"),
                  ("浮点", "Float32/Float64、精确值与容差", "NaN 策略、区间/相对误差"),
                  ("文本/字节", "UTF-8、UTF-16LE、字节数组、通配模式", "编码探测、结构模板"),
                  ("字节序", "Little / Big，由地址域默认并可覆盖", "Mixed-endian 特殊规则"),
                  ("再扫条件", "等于、不等、增加、减少、变化、未变化、区间", "未知初值、百分比变化、历史表达式"),
              ], [1600, 4200, 3560])

    add_heading(doc, "9.3 结果存储", 2)
    add_para(doc, "初扫不把所有候选对象放进托管堆。建议将候选地址编码为按区域排序的 64 位差分流，旧值放入独立列式块；块带校验和、目标进程实例、数据类型、字节序和扫描条件。再扫顺序读取旧块并生成新快照，成功后原子切换。")
    add_code(doc, "ScanSession\n  session.json           // 目标、适配器、类型、条件、版本\n  snapshot-0001.idx      // 区域目录与块偏移\n  snapshot-0001.addr     // UInt64 delta-coded addresses\n  snapshot-0001.value    // previous values\n  snapshot-0001.crc      // 块校验\n  audit.ndjson           // 读写失败与用户操作")

    add_heading(doc, "9.4 写入、冻结与条件规则", 2)
    for text in [
        "一次性写入：先读回并显示当前值、目标值、地址域、解析证据；用户确认后写入并校验实际写入字节数。",
        "批量写入：逐项返回成功/失败，不把部分成功伪装成事务；可选“全部预检通过后再写”。",
        "冻结：由单一调度器管理，不为每个地址创建线程；按组设置 16–1000ms 周期、抖动和最大写频率。",
        "条件写入：比较读取值后再写，支持 expected-old-value 以避免对变化中的地址误写。",
        "停止与回滚：提供全局紧急停止；只有用户明确选择且进程实例未变化时才恢复原值。",
        "所有写操作写入本地审计日志，包含时间、地址表达式、解析 Host 地址、旧值、新值和适配器版本。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "10. 模拟器适配器协议", 1)
    add_callout(doc, "核心原则", "扫描引擎只认识“内存域”和“已解析的 Host 地址”；模拟器版本、Guest 地址、镜像和字节序全部由适配器负责。")
    add_table(doc,
              ["接口", "输入/输出", "要求"],
              [
                  ("Probe", "进程列表 → 匹配分数、版本、架构", "不得只按 exe 名；至少结合路径、文件版本、模块或签名。"),
                  ("Attach", "ProcessInstance → AdapterSession", "返回能力集和诊断信息；不在 Probe 阶段申请写权限。"),
                  ("EnumerateDomains", "→ RAM/VRAM/SRAM/ROM 等域", "每个域声明大小、地址宽度、字节序、可写性和生命周期。"),
                  ("Resolve", "LogicalAddress → ResolvedAddress", "返回证明链；不确定时失败，禁止猜测。"),
                  ("Read/Write", "域地址 + 长度/数据", "优先走模拟器原生 API；否则委托 Broker 的 Host 后端。"),
                  ("Pause/Resume", "可选", "只有能力协商通过才显示 UI 按钮。"),
                  ("HealthCheck", "→ Attached/Reloaded/Stale/Failed", "ROM 切换、存档载入或映射变化后使旧解析失效。"),
                  ("CreateDiagnostics", "→ 脱敏诊断包", "含版本、模块、内存域、签名命中和错误，不含 ROM/内存正文。"),
              ], [1750, 2880, 4730], font_size=9.1)

    add_heading(doc, "10.1 地址解析策略优先级", 2)
    strategies = [
        "模拟器公开调试 API / IPC / 插件 API：语义最稳定，可直接获得内存域、暂停和帧同步。",
        "Libretro 桥接：由受控前端或桥接组件取得 RETRO_MEMORY_SYSTEM_RAM 等域；不能假定任意 RetroArch 外部进程都直接暴露这些指针。",
        "模块基址 + 导出符号或稳定偏移：仅在版本严格匹配时启用。",
        "签名定位 + 结构验证：签名命中后必须验证区域大小、权限、魔数或已知内存行为。",
        "用户校准的手工基址：只用于开发/诊断，不作为默认发布配置。",
    ]
    for text in strategies:
        add_number(doc, text, strategy_num_id)

    add_heading(doc, "10.2 适配器清单示例", 2)
    add_code(doc, "schemaVersion: 1\nid: org.fpenext.adapter.example\ndisplayName: Example Emulator\nsupported:\n  - process: ExampleEmu.exe\n    arch: [x64]\n    versions: ['1.2.*']\ncapabilities: [read, write, memoryDomains, pause]\ndomains:\n  - id: console.ram\n    guestBase: 0x00000000\n    size: 0x02000000\n    endian: little\nresolver:\n  kind: signature\n  module: ExampleEmu.exe\n  patternId: ram-base-v1\n  validation: [committed, readable, expected-size]\n")

    add_heading(doc, "10.3 版本与发布策略", 2)
    for text in [
        "适配器包独立语义化版本；清单声明核心协议版本范围。",
        "默认拒绝未签名适配器；开发模式需显式开启并显示持续警告。",
        "同一模拟器的不同主版本可以共享代码，但必须各自拥有 Probe、Resolver 和回归夹具。",
        "发现未知版本时进入只读诊断模式，不复用“最接近版本”的写入偏移。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "11. 逐模拟器调试流程", 1)
    steps = [
        "锁定对象：记录模拟器安装包来源、版本、架构、可执行文件 SHA-256、核心/插件版本和测试游戏。",
        "建立基准：选择一个可稳定观察的数值或内存片段，分别记录 Guest 地址与游戏内行为。",
        "确认能力：优先检查官方调试 API、IPC、脚本、插件或 Libretro 内存域；没有接口才走进程内存。",
        "定位 Host 映射：找出基址、区域大小、权限、字节序、镜像、分段或 bank switching。",
        "做生命周期实验：重启模拟器、重新载入 ROM、载入存档、切换渲染器/核心，观察映射是否变化。",
        "实现 Probe 与 Resolve：未知版本必须失败；解析结果包含可复核证据。",
        "生成夹具：保存脱敏的区域元数据、签名命中、地址翻译样例和合成内存块，不保存受版权保护 ROM。",
        "验证读：对已知地址连续读取，检查帧间变化、越界、镜像和端序。",
        "验证写：仅在离线副本中执行小范围、可恢复写入；验证错误地址不会被触及。",
        "发布候选：跑核心回归、版本矩阵、未知版本拒绝测试，最后签名适配器包。",
    ]
    for text in steps:
        add_number(doc, text, steps_num_id)

    add_heading(doc, "11.1 每个模拟器的交付物", 2)
    add_table(doc,
              ["交付物", "内容", "退出条件"],
              [
                  ("support.yaml", "支持版本、架构、进程标识、能力", "未知版本不会被误识别。"),
                  ("resolver", "内存域与 Guest→Host 翻译", "重启与载入测试仍能解析或明确失效。"),
                  ("fixtures", "合成内存、区域表、签名与期望结果", "CI 可离线复现。"),
                  ("diagnostics", "探测日志与错误说明", "用户可一键导出，不包含内存正文。"),
                  ("compatibility.md", "已测版本、限制、写入风险", "UI 能显示同样的信息。"),
              ], [1850, 4350, 3160])

    add_heading(doc, "11.2 推荐适配顺序", 2)
    add_table(doc,
              ["阶段", "对象", "理由"],
              [
                  ("A", "通用 Win32/x64 进程后端 + 合成靶程序", "先证明高地址、权限、区域变化和扫描正确性。"),
                  ("B", "Legacy INI 兼容适配器", "读取六条旧记录，验证旧地址转换语义；不承诺老程序本身在 Win11 运行。"),
                  ("C", "Libretro 桥接", "统一接口可覆盖多个核心，但需要受控前端/桥接层与逐核心验证。"),
                  ("D", "用户选定的第一个现代独立模拟器", "一次只解决一个版本族，建立团队适配模板。"),
                  ("E", "后续模拟器", "复用协议、诊断与夹具，按需求排序。"),
              ], [1100, 3300, 4960])

    add_heading(doc, "12. 前端信息架构与旧功能迁移", 1)
    add_table(doc,
              ["新页面", "核心任务", "旧功能来源"],
              [
                  ("连接", "选择进程/模拟器、架构、权限、版本、适配器健康", "TargetEdit / GameBitBtn / Fpe.ini"),
                  ("扫描", "数据类型、端序、初扫/再扫、条件、进度、取消", "TabScan / Mission / FPEMEM*.DAT"),
                  ("结果", "64 位地址、Guest/Host 切换、当前/旧值、区域、固定与导出", "AddrList / ListButton / EditButton"),
                  ("地址表", "规则、条件、冻结组、备注、启停、审计", "Address / Lock Value / if / Comment"),
                  ("内存查看器", "虚拟化十六进制、文本、跳转、只读/编辑提交", "TabEditor / SGrid / Hex/String"),
                  ("适配器", "支持版本、能力、诊断、开发模式、校准向导", "TabEmu / Emulator INI"),
                  ("宏与输入", "前台热键、动作序列、紧急停止", "TabMacKey / tt.dll"),
                  ("设置与诊断", "临时目录、性能、日志、权限、版本", "TabOthers / Registry / About"),
              ], [1600, 4300, 3460], font_size=9.1)

    add_heading(doc, "12.1 关键 UX 约束", 2)
    for text in [
        "顶部连接条始终显示：进程名、PID、启动时间、x86/x64、适配器版本、只读/可写状态。",
        "Guest 与 Host 地址使用不同颜色与明确标签；复制菜单要求用户选择地址空间。",
        "扫描条件以状态机呈现，禁止在没有兼容旧快照时继续再扫。",
        "写入/冻结操作使用危险色、二次确认和全局停止；只读模式下不渲染可误触的写按钮。",
        "百万级结果必须分页/虚拟化；筛选、排序和固定不复制全部对象。",
        "错误提示包含地址、区域、Win32 错误、建议动作和“生成诊断包”入口。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "13. 旧数据迁移", 1)
    add_table(doc,
              ["旧数据", "迁移策略", "风险控制"],
              [
                  ("Fpe.ini", "只读解析 Count/E##；原字符串和十六进制参数原样保存；用户确认后绑定新适配器。", "不自动把参数解释为基址/偏移。"),
                  (".fpe 地址表", "收集真实样例后做格式侦测；第一版导入到草稿并逐项校验。", "未知字段保留 raw blob；不覆盖原文件。"),
                  (".mis/.mi0", "仅在格式清晰且目标进程实例匹配时转换。", "默认不恢复陈旧扫描快照。"),
                  ("FPEMEM*.DAT", "视为临时搜索索引，不作为长期兼容目标。", "可提供一次性离线分析器，但不直接继续扫描。"),
                  ("注册表 SOFTWARE\\jaw\\fpe", "仅在用户选择“导入旧设置”时读取；迁入 SQLite。", "不写回旧键，不要求管理员。"),
              ], [1700, 4780, 2880], font_size=9.2)

    add_heading(doc, "14. 安全与运行约束", 1)
    for text in [
        "默认只读：连接和扫描不申请写权限。开启写入后，UI 显示持续状态条并允许立即撤销授权。",
        "最小权限：只申请当前操作需要的 PROCESS_* 权限；不默认启用 SeDebugPrivilege。",
        "目标确认：写入前验证 PID、启动时间、映像路径、适配器会话和解析代次。",
        "适配器隔离：禁止任意本地 DLL 侧载；发布适配器签名校验，开发适配器独立目录和明显标记。",
        "数据隐私：诊断包默认不含内存内容、ROM 路径中的用户名或完整文件列表；导出前预览。",
        "应用范围：仅面向本地离线模拟器与用户有权调试的程序；检测到受保护进程或反作弊组件时拒绝附加。",
        "供应链：主程序、Broker、适配器和更新清单均签名；生成 SBOM，固定依赖版本并定期升级。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "15. 测试策略", 1)
    add_heading(doc, "15.1 必建合成靶程序", 2)
    add_para(doc, "构建 x86 与 x64 两个测试目标，按命令创建可读、只读、不可访问、Guard、Mapped、Image 和动态释放区域；x64 靶明确在 0x0000000100000000 以上分配/暴露可验证值。测试不依赖任何商业游戏。")
    add_table(doc,
              ["层级", "覆盖", "关键断言"],
              [
                  ("单元", "地址解析、端序、类型比较、差分编码、条件规则", "边界值、溢出、NaN、未对齐、UInt64 不截断。"),
                  ("互操作", "OpenProcess、VirtualQueryEx、Read/WriteProcessMemory", "权限最小化、部分失败、句柄释放、正确错误码。"),
                  ("集成", "x86/x64 靶、区域变化、进程重启、PID 复用", "旧会话失效；高地址读写正确。"),
                  ("适配器", "Probe、域、Resolve、未知版本、重载", "不猜测、不跨版本写入、诊断完整。"),
                  ("性能", "1/4/16GiB 可读空间、稀疏区域、百万结果", "UI 始终响应；取消延迟 < 200ms；内存受控。"),
                  ("UI", "高 DPI、中文/英文、键盘、无障碍、危险操作", "地址不截断；焦点清晰；只读模式不可误写。"),
              ], [1400, 3930, 4030], font_size=9.2)

    add_heading(doc, "15.2 首版验收指标", 2)
    add_table(doc,
              ["指标", "目标", "说明"],
              [
                  ("地址正确性", "100% 通过 >4GB 合成地址用例", "含扫描、显示、复制、保存、重载、写入。"),
                  ("读写安全", "任何未知版本或会话过期均拒绝写", "失败必须有可操作诊断。"),
                  ("取消响应", "用户取消后 200ms 内停止提交新结果", "底层系统调用返回后清理。"),
                  ("扫描吞吐", "参考机 Int32 精确初扫 ≥ 500 MiB/s", "作为调优目标，最终以指定硬件基线复测。"),
                  ("UI 内存", "百万结果不生成百万 ViewModel", "依赖分页/虚拟化和磁盘快照。"),
                  ("稳定性", "8 小时冻结/监控无句柄泄漏和未解释写失败", "使用合成靶和离线模拟器。"),
              ], [1800, 3370, 4190])

    add_heading(doc, "16. 实施路线图", 1)
    add_table(doc,
              ["阶段", "范围", "退出条件", "粗略工作量"],
              [
                  ("P0 规格冻结", "样本归档、风险边界、数据模型、测试靶设计", "ADR 与接口评审通过", "1 周"),
                  ("P1 64 位核心", "x64 Broker、区域枚举、读、初扫/再扫、磁盘快照", ">4GB 全链路测试通过", "3–4 周"),
                  ("P2 最小前端", "连接、扫描、结果、内存查看器、诊断", "可完成只读扫描工作流", "2–3 周"),
                  ("P3 写入与地址表", "一次写、冻结组、条件、审计、紧急停止", "安全用例与 8 小时稳定性通过", "2–3 周"),
                  ("P4 适配器 SDK", "协议、隔离主机、清单、夹具、诊断包", "示例适配器与未知版本拒绝通过", "2–3 周"),
                  ("P5 第一个模拟器", "按用户选择完成一个版本族", "读写/重启/载入/版本矩阵通过", "1–4 周/模拟器"),
                  ("P6 迁移与发布", "旧 INI/.fpe 导入、签名、安装、文档", "可回滚发布与数据不覆盖", "2 周"),
              ], [1250, 3300, 3190, 1620], font_size=9.0)
    add_para(doc, "单人从零完成可用 MVP 约 10–14 个开发周；首个模拟器适配通常再需 1–4 周。差异主要来自是否存在稳定官方接口、内存映射是否跨版本变化，以及需要支持多少核心/渲染后端。", italic=True, color=MUTED)

    add_heading(doc, "17. Definition of Done", 1)
    dod = [
        "FpeNext.UI、Broker、AdapterHost 均为 x64，可在受支持的 Windows 11 x64 环境安装和卸载。",
        "可附加 x86/x64 合成靶，正确枚举区域并读取 0xFFFFFFFF 以上地址。",
        "Int8–64、UInt8–64、Float32/64、字节数组、UTF-8/UTF-16 和大小端扫描通过测试。",
        "结果快照可取消、恢复、校验；进程重启后旧地址自动失效。",
        "一次写入、冻结、条件写入、全局停止与审计日志通过安全测试。",
        "适配器未知版本拒绝写，诊断包不泄露 ROM 或内存正文。",
        "Legacy Fpe.ini 只读导入并保留原始字段；原文件不修改。",
        "首个用户选定模拟器完成明确版本矩阵和回归夹具。",
        "所有发布二进制签名，SBOM 和第三方许可证齐备。",
    ]
    for text in dod:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "18. 下一次调试一个模拟器时需要的输入", 1)
    add_callout(doc, "最小输入包", "模拟器精确版本与下载来源、exe SHA-256、x86/x64、测试游戏/ROM 的合法副本、一个可观测数值、Guest 地址线索、是否允许离线写入。")
    for text in [
        "模拟器完整版本号、安装模式（便携/安装）、核心或插件版本。",
        "目标游戏、区域版本、存档状态和复现步骤。",
        "希望优先实现的能力：只读扫描、写入、冻结、暂停、帧同步、指针链。",
        "已知 Cheat 地址、调试器截图或内存域文档（若有）。",
        "允许的测试边界：仅本地离线；是否可使用模拟器自带调试接口。",
    ]:
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "附录 A：样本哈希与 PE 摘要", 1)
    add_table(doc,
              ["对象", "SHA-256", "摘要"],
              [
                  ("ZIP", "0B48C37C6C8A7D5049EEFF8311195D0DCE71E6A2E8889D6D27E718ED94B341E0", "用户提供压缩包"),
                  ("FPE.exe", "9225B6D336E474E2534D20E2FBDF3127D808308BDE9665CBB2AB735A62B0E738", "PE32 x86；无签名；ASLR/NX 关闭"),
                  ("SPE.exe", "2C4CDA2F08660CFC742F4A651241588D87A21E68AF483BF8F24C2033FB3D7AB7", "PE32 x86；无签名；图片工具"),
                  ("tt.dll", "8C621D4D9CD26A72CF9081EF7B036D4256EEC33197FA7F80EAF34569F8D5CEE3", "PE32 x86；导出 Setup/UnSetup/Setkey"),
                  ("Fpe.ini", "7FF73C36E53FA3BE9AAC162994F491ACF20509C222B2705B203F943CE1106F6A", "六条模拟器映射"),
              ], [1450, 5650, 2260], font_size=8.3)
    add_para(doc, "注：PE 时间戳相互矛盾（FPE.exe 甚至指向 2037），不能作为可信构建日期。", italic=True, color=MUTED)

    add_heading(doc, "附录 B：关键静态证据", 1)
    evidence = [
        ("FPE 导入", "OpenProcess、VirtualQueryEx、ReadProcessMemory、WriteProcessMemory、CreateThread、RegisterHotKey"),
        ("FPE 字符串", "FPEMEM%02d.DAT、Scan:(%x-%x)、.fpe、.mis、%08X、>4GB、Can't calculate address"),
        ("主窗体", "地址表列 Address / Lock Value / if / Comment；Memory；Data Type 8/16/32；LockTrackBar"),
        ("模拟器配置窗体", "PS, SFC 标签；Load/Save/Translate；FPE.INI/Emulator/Count/E%02d"),
        ("tt.dll", "SetWindowsHookExA、keybd_event、mouse_event、共享节、Setup/UnSetup/Setkey"),
        ("SPE", "Show Picture Expert、BMP/JPG/GIF、JPG quality、Zoom、图片文件列表"),
    ]
    add_table(doc, ["来源", "证据"], evidence, [1800, 7560])

    add_heading(doc, "附录 C：外部技术参考", 1)
    refs = [
        ("Microsoft Learn：ReadProcessMemory", "https://learn.microsoft.com/en-us/windows/win32/api/memoryapi/nf-memoryapi-readprocessmemory"),
        ("Microsoft Learn：WriteProcessMemory", "https://learn.microsoft.com/en-us/windows/win32/api/memoryapi/nf-memoryapi-writeprocessmemory"),
        ("Microsoft Learn：VirtualQueryEx", "https://learn.microsoft.com/en-us/windows/win32/api/memoryapi/nf-memoryapi-virtualqueryex"),
        ("Microsoft Learn：IsWow64Process2", "https://learn.microsoft.com/en-us/windows/win32/api/wow64apiset/nf-wow64apiset-iswow64process2"),
        ("Microsoft Learn：Process Security and Access Rights", "https://learn.microsoft.com/en-us/windows/win32/procthread/process-security-and-access-rights"),
        ("Microsoft Learn：.NET Native Interop Best Practices", "https://learn.microsoft.com/en-us/dotnet/standard/native-interop/best-practices"),
        ("Microsoft：.NET Support Policy", "https://dotnet.microsoft.com/en-us/platform/support/policy"),
        ("Libretro：Core Development Overview", "https://docs.libretro.com/development/cores/developing-cores/"),
        ("Libretro：canonical libretro.h", "https://github.com/libretro/libretro-common/blob/master/include/libretro.h"),
    ]
    for title, url in refs:
        p = doc.add_paragraph()
        apply_num(p, bullet_num_id, 0)
        p.paragraph_format.space_after = Pt(4)
        add_hyperlink(p, title, url)
    add_para(doc, "资料用途：Win32 API 权限/读写语义、目标架构识别、.NET 互操作与 LTS 选择，以及 Libretro 内存域/前端-核心接口设计。", italic=True, color=MUTED)

    add_heading(doc, "附录 D：仍需项目方决策", 1)
    decisions = [
        "首个正式支持的现代模拟器及精确版本族。",
        "是否只做只读扫描 MVP，还是第一版同时包含写入/冻结。",
        "是否需要兼容现有 .fpe 文件；如需要，提供至少 3–5 个真实样例和对应界面含义。",
        "发布方式：便携包、传统安装器或 MSIX；是否需要自动更新。",
        "宏与 SPE/GPE 图片工具是否仍有实际用户需求。",
        "项目许可证、适配器分发政策与代码签名证书归属。",
    ]
    for text in decisions:
        add_bullet(doc, text, bullet_num_id)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    bullet_num_id, number_num_ids = add_num_definitions(doc)
    configure_header_footer(doc)
    doc.core_properties.title = "FPE 2001 Win11 重构分析与架构设计"
    doc.core_properties.subject = "64 位进程内存扫描与多模拟器适配器架构"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "FPE2001, Windows 11, x64, emulator adapter, memory scanner"
    doc.core_properties.comments = "基于用户提供压缩包的只读静态分析；未执行样本。"
    add_cover(doc)
    add_contents(doc, number_num_ids[0])
    add_body(doc, bullet_num_id, number_num_ids)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    main()
