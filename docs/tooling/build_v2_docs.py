#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import build_refactor_doc as base
import build_implementation_specs as impl


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ARCH_OUT = OUT / "FPE2001-Remake_Win11重构分析与架构设计_v2.1.docx"
CODE_OUT = OUT / "FPE2001-Remake_代码实施规格包_v2.1.docx"
UI_OUT = OUT / "FPE2001-Remake_UIUX实施规格_v2.1.docx"


MODULES = [
    ("M01", "扫描", "TabScan", "进程/模拟器选择、任务 Mission、首次与再次扫描、8/16/32/未知类型、候选列表、添加/编辑。"),
    ("M02", "地址表", "TabTabs", "地址条目、立即写入、修改、锁定/条件、注释、游戏信息、加载/保存 .fpe。"),
    ("M03", "十六进制编辑", "TabEditor", "打开数据源、查找、撤销、标签/注释、刷新、16 进制与字符串显示；覆盖式编辑。"),
    ("M04", "文件", "TabSelect", "驱动器/目录浏览、子目录列表、全部文件、刷新、扫描、停止、编辑/保存。"),
    ("M05", "图片", "TabGPE + SPE", "图片目录、预览、命名、删除、编辑、连续编号、x2；SPE 支持 BMP/JPG/GIF 与缩放/导出。"),
    ("M06", "宏与热键", "TabMacKey + tt.dll", "设置/删除/更新宏，键盘与鼠标左/右/双击动作，热键与前台窗口控制。"),
    ("M07", "速度", "TabSpeed", "开关与滑杆控制速度；新版仅在目标/适配器声明能力时启用。"),
    ("M08", "设置与工具", "TabOthers", "内存上限、临时目录、SPE、撤销、锁定间隔、字符串列表与工具参数。"),
    ("M09", "关于", "TabAbout", "版本、作者/来源、许可证、组件与诊断入口。"),
]


def p(doc, text, **kwargs):
    return base.add_para(doc, text, **kwargs)


def h(doc, text, level=1):
    return base.add_heading(doc, text, level)


def bullets(doc, items, num_id, level=0):
    for item in items:
        base.add_bullet(doc, item, num_id, level)


def numbers(doc, items, num_id):
    for item in items:
        base.add_number(doc, item, num_id)


def code(doc, text):
    return base.add_code(doc, text.strip("\n"))


def cover(doc, kicker, title, subtitle, rows, lead):
    p(doc, kicker, size=10, bold=True, color=base.GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, before=58, after=22)
    p(doc, "FPE 2001 → FPE2001-Remake", size=30, bold=True, color=base.INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p(doc, title, size=22, bold=True, color=base.BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    p(doc, subtitle, size=13, color=base.MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=48)
    table = doc.add_table(rows=len(rows), cols=2)
    base.set_table_geometry(table, [2300, 7060], indent=120)
    base.set_table_borders(table, color="D8DEE5", size=5)
    for i, (label, value) in enumerate(rows):
        base.set_cell_shading(table.cell(i, 0), base.LIGHT_BLUE)
        base.set_run_font(table.cell(i, 0).paragraphs[0].add_run(label), base.FONT, 10, bold=True, color=base.INK)
        base.set_run_font(table.cell(i, 1).paragraphs[0].add_run(value), base.FONT, 10, color="20262D")
    p(doc, lead, size=11.3, bold=True, color=base.DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=36, after=7)
    p(doc, "v2.1 以原版九模块和原始工作流为功能基线，现代化仅改变实现、安全边界与交互表达。", size=9.6, color=base.MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def finish(doc, path):
    impl.trim_trailing_empty_paragraphs(doc)
    impl.add_terminal_paragraph(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def new_doc(short_title):
    doc = Document()
    base.configure_styles(doc)
    bullet, decimals = base.add_num_definitions(doc)
    impl.set_header_footer(doc, short_title)
    return doc, bullet, decimals


def module_table(doc):
    base.add_table(doc, ["ID", "新版模块", "原版证据", "必须保留的功能"], MODULES,
                   [850, 1450, 1700, 5360], font_size=8.8)


def build_architecture():
    doc, bullet, dec = new_doc("FPE2001-Remake Win11 重构设计 v2.1")
    cover(doc, "RECONSTRUCTION ANALYSIS & ARCHITECTURE", "Win11 重构分析与架构设计 v2.1",
          "完整保留原版九模块、文件十六进制编辑与模拟器逐个适配",
          [("分析对象", "FPE2001_绿色版.zip；静态分析，不运行未知二进制"),
           ("功能基线", "TabScan / TabTabs / TabEditor / TabSelect / TabGPE / TabMacKey / TabSpeed / TabOthers / TabAbout"),
           ("目标平台", "Windows 11 x64；超 4GB 地址；大文件偏移；多模拟器适配器"),
           ("版本", "v2.1 · 2026-08-16")],
          "结论：保留原功能表面和模块顺序，底层整体重写；文件编辑与进程内存编辑必须共享同一十六进制工作台。")
    impl.add_contents(doc, ["结论与设计原则", "静态分析证据与边界", "原版九模块功能基线", "原版关键工作流",
                            "v2.1 目标与非目标", "目标架构", "统一字节源与十六进制编辑", "文件浏览与内容扫描",
                            "进程扫描、地址表与冻结", "图片、宏、速度和工具", "模拟器适配器平台", "64 位地址与文件偏移",
                            "安全、兼容与数据迁移", "UI 信息架构", "实施路线图与验收", "风险与待定项"], dec[0])

    h(doc, "1. 结论与设计原则")
    base.add_callout(doc, "核心决定", "不再把原版“编辑器”和“文件”功能降级为附属工具。新版以九模块为一等功能，并让十六进制工作台同时处理进程内存与本地文件。")
    bullets(doc, [
        "功能等价优先：原版按钮或流程能映射到明确的新入口；无法确认语义的字段保留为兼容元数据，不猜测。",
        "原设计顺序保留：扫描→地址表→编辑→文件→图片→宏→速度→设置→关于，降低老用户迁移成本。",
        "现代化不等于删功能：文件浏览/扫描、GPE/SPE、宏和速度能力均进入路线图，只按风险分阶段交付。",
        "共享核心：进程内存与文件都抽象为 64 位可寻址字节源，复用网格、搜索、编码、书签、撤销和差异显示。",
        "安全边界：本地离线单机；不做内核驱动、DLL 注入、反作弊绕过、隐藏或联网对战支持。",
    ], bullet)

    h(doc, "2. 静态分析证据与边界")
    base.add_table(doc, ["对象", "静态证据", "可得结论"], [
        ("FPE.exe", "PE32 x86、VCL DFM、OpenProcess/VirtualQueryEx/ReadProcessMemory/WriteProcessMemory、九个 TabSheet", "核心为进程扫描/写入工具，UI 功能可从组件与事件名还原。"),
        ("Fpe.ini", "6 条 Emulator 记录，内部格式字符串 %d,%X,%X,%s", "存在旧模拟器地址映射；字段精确语义仍需动态验证。"),
        ("SPE.exe", "Show Picture Expert；BMP/JPG/GIF、缩放、JPG 质量", "是图片浏览/转换辅助程序，不属于内存扫描后端。"),
        ("tt.dll", "SetWindowsHookExA、keybd_event、mouse_event、前台窗口 API", "为宏/热键支持；新版应减少全局钩子依赖。"),
        ("限制", "不运行样本、不注入、不抓动态内存", "文件格式、速度实现与部分按钮语义需要兼容测试确认。"),
    ], [1500, 4300, 3560], font_size=8.9)
    p(doc, "样本 ZIP SHA-256：0B48C37C6C8A7D5049EEFF8311195D0DCE71E6A2E8889D6D27E718ED94B341E0。所有可执行文件均未签名，FPE/SPE/tt 为 32 位。", size=9.4, color=base.MUTED)

    h(doc, "3. 原版九模块功能基线")
    module_table(doc)
    h(doc, "3.1 M03 十六进制编辑器的原始线索", 2)
    bullets(doc, [
        "顶部动作：Open(F1)、Find(F2)、Undo(F3)、Add Tab(F4)、Fresh(F5)、Note(F6)。",
        "中央为 TColorStringGrid；底部含覆盖编辑框、16 进制/10 进制或字符串模式切换。",
        "与 M04 文件页的 Edit/Save(F5) 联动，说明本地文件能够送入编辑器读取和修改。",
        "新版必须明确区分 File、ProcessMemory、Snapshot 三种来源，并显示未保存修改。",
    ], bullet)
    h(doc, "3.2 M04 文件页的原始线索", 2)
    bullets(doc, [
        "DriveComboBox、DirectoryListBox、FilterComboBox 和文件 ListView 构成完整文件浏览器。",
        "动作包含列出子目录(F1)、全部文件(F2)、刷新(F3)、扫描(F4)、编辑/保存(F5)、停止访问(F6)。",
        "新版使用系统路径语义与长路径支持，但保留目录树、筛选、递归扫描、取消和送入编辑器。",
    ], bullet)

    h(doc, "4. 原版关键工作流")
    base.add_table(doc, ["工作流", "原版路径", "v2.1 保留方式"], [
        ("找值并锁定", "TabScan 扫描→AddrList→TabTabs 写入/锁定", "扫描会话→候选→地址表→条件冻结；保留任务/Mission。"),
        ("浏览并改文件", "TabSelect 浏览/扫描→EditSave→TabEditor", "文件页→以文件源打开十六进制工作台→差异→备份/原子保存。"),
        ("直接看内存", "TabEditor 打开/刷新→网格→覆盖编辑", "以进程内存源打开；读回校验、权限提示、刷新代次。"),
        ("图片采集", "TabGPE 目录/命名/编号→SPE 查看/转换", "图片工作区+可选外部查看器；元数据与核心扫描器解耦。"),
        ("自动化", "TabMacKey→tt.dll→键鼠/窗口", "显式热键、动作列表、前台目标校验、紧急停止。"),
        ("速度", "TabSpeed 开关+滑杆", "适配器能力；无能力时页面保留但解释不可用。"),
    ], [1800, 3100, 4460], font_size=8.9)

    h(doc, "5. v2.1 目标与非目标")
    h(doc, "5.1 必须实现", 2)
    bullets(doc, [
        "Windows 11 x64、x86/x64 目标、宿主地址超过 4GB、文件偏移超过 4GB。",
        "九模块壳层和快捷键映射；原版 Mission、.fpe、.mis 等提供受控导入/导出或兼容占位。",
        "可取消扫描、磁盘结果快照、地址表、写入、条件冻结、回滚和全局停止。",
        "本地文件浏览、内容扫描、十六进制/文本搜索、覆盖/插入、编码切换、撤销与安全保存。",
        "图片、宏、速度与设置具有独立服务接口和能力探测，不硬连扫描核心。",
    ], bullet)
    h(doc, "5.2 明确非目标", 2)
    bullets(doc, ["不运行或复用旧版 DLL；不修补旧 EXE。", "不承诺所有模拟器一次支持；按精确版本逐个适配。", "不默认递归扫描整个磁盘；不自动修改受保护目录。", "不将宏用于绕过安全或控制非用户选择的应用。"], bullet)

    h(doc, "6. 目标架构")
    code(doc, "Fpe2001Remake.UI (九模块壳层)\n  ├─ Application / Domain / Contracts\n  ├─ Memory.Win32 + Scan + AddressBook\n  ├─ BinaryEditor.Core + ByteSources.File + ByteSources.Process\n  ├─ FileWorkspace + MediaWorkspace + Automation + SpeedControl\n  ├─ Storage + LegacyImport\n  └─ IPC → AdapterHost / optional Broker")
    base.add_table(doc, ["层/进程", "职责", "依赖约束"], [
        ("UI", "九模块导航、页面状态、快捷键、对话框", "不得直接 P/Invoke、直接打开 SQLite 或持有文件句柄。"),
        ("Application", "会话、命令编排、跨模块跳转", "依赖抽象；不依赖 WPF View。"),
        ("Core services", "扫描、编辑、文件搜索、图片、宏、速度", "共享 Domain/Contracts；模块之间通过用例而非 UI 引用。"),
        ("AdapterHost", "模拟器探测、地址翻译、速度/暂停等能力", "一实例一适配器；崩溃不拖垮 UI。"),
        ("Broker", "按需最小权限读写", "只接受白名单 PID、范围与命令。"),
    ], [1700, 3600, 4060], font_size=9)

    h(doc, "7. 统一字节源与十六进制编辑")
    base.add_callout(doc, "关键抽象", "十六进制网格不关心数据来自文件还是进程。IByteSource 负责 64 位读取；可编辑来源额外实现变更集和提交策略。")
    base.add_table(doc, ["来源", "长度/寻址", "编辑语义", "提交"], [
        ("File", "ulong 文件长度与偏移", "覆盖、插入、删除；稀疏变更集", "外部变化检测→备份→临时文件→原子替换。"),
        ("ProcessMemory", "逻辑/宿主地址；区域可能不可读", "默认覆盖；不可插入/删除", "写前比较→写入→读回；进程身份变化拒绝。"),
        ("Snapshot", "只读快照", "只读或导出补丁", "另存为文件，不写回原目标。"),
    ], [1700, 2400, 2600, 2660], font_size=9)
    bullets(doc, [
        "网格：偏移、16 字节、文本区；编码 ASCII/UTF-8/UTF-16LE/GBK；地址/偏移使用等宽字体。",
        "搜索：十六进制（支持 ??）、文本、整数、浮点；范围、方向、大小写、端序可选。",
        "编辑：覆盖/插入、撤销/重做、书签、注释、修改高亮、差异预览；大文件按页缓存。",
        "安全：默认只读打开；保存前确认外部修改、权限、备份路径与最终 SHA-256。",
    ], bullet)

    h(doc, "8. 文件浏览与内容扫描")
    bullets(doc, [
        "文件页保留驱动器/目录/筛选/子目录/全部文件/刷新/扫描/停止/编辑保存的原工作流。",
        "默认不跟随目录联接和符号链接；访问拒绝、长路径、文件消失和并发修改为正常状态。",
        "内容扫描支持十六进制、文本和数值模式；结果记录路径、偏移、预览、匹配类型和文件时间。",
        "点击结果以只读方式定位到十六进制编辑器；用户显式启用编辑后才创建变更集。",
        "大目录扫描显示已访问文件数、字节数、匹配数、跳过数和取消；支持扩展名、大小、时间过滤。",
    ], bullet)

    h(doc, "9. 进程扫描、地址表与冻结")
    base.add_table(doc, ["子系统", "v2.1 要求"], [
        ("区域枚举", "VirtualQueryEx；MEM_COMMIT；跳过 NOACCESS/GUARD；checked 地址推进。"),
        ("扫描", "整数/浮点/字节/文本、端序、未知值、首次/再次、任务 Mission、取消。"),
        ("结果", "磁盘快照、窗口化读取、候选历史、从结果送入地址表或编辑器。"),
        ("地址表", "逻辑地址优先、宿主证据、标签/注释、立即写入、条件冻结、.fpe 兼容。"),
        ("冻结", "单调度器、最短 16ms、默认 250ms、连续失败停用、500ms 全停。"),
    ], [1900, 7460], font_size=9.2)

    h(doc, "10. 图片、宏、速度和工具")
    base.add_table(doc, ["模块", "保留", "现代化约束"], [
        ("图片", "目录、预览、命名、删除、连续编号、x2、SPE 格式转换", "使用受支持图像解码；不把图片解析放进 UI 线程；保留 EXIF/质量选项。"),
        ("宏", "设置/删除/更新、键盘、鼠标左/右/双击、热键", "优先 RegisterHotKey；动作绑定当前目标；有急停、倒计时与可见运行状态。"),
        ("速度", "开关、滑杆", "只通过适配器官方/可验证能力；无能力时禁用；不做通用时间钩子。"),
        ("设置", "内存上限、临时目录、SPE、撤销、锁定间隔、字符串列表", "拆为扫描/存储/写入/工具/隐私/开发者分组；提供旧配置导入预览。"),
    ], [1500, 3000, 4860], font_size=8.9)

    h(doc, "11. 模拟器适配器平台")
    numbers(doc, ["官方调试 API、IPC 或插件。", "Libretro 内存域桥接。", "模块导出或稳定偏移，并严格匹配版本。", "签名定位加结构验证。", "手动校准只用于开发者模式。"], dec[1])
    bullets(doc, [
        "能力包括 Read、Write、Domains、GuestVirtualTranslation、PauseResume、FrameSync、SpeedControl、SaveRam、Diagnostics。",
        "速度页和部分图片/捕获能力由适配器能力驱动；页面不因能力缺失而消失。",
        "每个解析地址包含 ResolutionProof；进程重启或目标版本变化后旧证明失效。",
    ], bullet)

    h(doc, "12. 64 位地址与文件偏移")
    base.add_table(doc, ["概念", "类型", "显示"], [
        ("HostVirtualAddress", "ulong/nuint", "0x + 16 位十六进制"),
        ("LogicalAddress", "AddressSpace + DomainId + ulong Value", "域名 + 适当宽度；可复制完整 JSON"),
        ("FileOffset", "ulong", "0x + 至少 16 位；同时可显示十进制"),
        ("Length", "ulong（领域）/ nuint（OS 边界）", "人类单位 + 精确字节"),
    ], [2300, 3600, 3460], font_size=9.2)
    base.add_callout(doc, "禁止", "核心领域模型不得使用 int/uint 存地址、偏移或文件长度；所有加法 checked；JSON 中的 ulong 使用十六进制字符串。")

    h(doc, "13. 安全、兼容与数据迁移")
    bullets(doc, [
        "文件编辑默认备份并原子保存；系统目录、只读文件和外部修改必须二次确认或拒绝。",
        "文件扫描默认当前目录，不默认全盘；递归、跟随链接和受保护目录均需显式选择。",
        "宏只对用户选择目标生效；焦点丢失、目标变化、急停和超时立即停止。",
        "旧 .fpe/.mis/.mi0/Fpe.ini 先预览再导入；未知字段放 legacy_metadata_json。",
        "诊断默认不包含完整内存、文件内容、宏按键序列或完整个人路径。",
    ], bullet)

    h(doc, "14. UI 信息架构")
    p(doc, "新版保留原版九页顺序，使用标题栏下方的横向模块带；窗口缩窄时模块带可水平滚动或进入“更多”，但顺序不变。每个模块顶部保留动作栏和 F 键提示。", color=base.MUTED)
    base.add_table(doc, ["顺序", "模块", "主要动作栏"], [
        ("1", "扫描", "清空、新任务、命名、加载/保存 Mission、计算、撤销、扫描/停止、添加/列表/编辑"),
        ("2", "地址表", "添加、删除、清空、修改、编辑、写入、加载/保存 .fpe"),
        ("3", "十六进制", "打开、查找、撤销、书签、刷新、注释、保存/另存"),
        ("4", "文件", "子目录、全部文件、刷新、扫描、编辑/保存、停止"),
        ("5-9", "图片/宏/速度/设置/关于", "按原模块动作保留；危险动作使用文字+图标。"),
    ], [900, 1800, 6660], font_size=9)

    h(doc, "15. 实施路线图与验收")
    base.add_table(doc, ["阶段", "范围", "出口标准"], [
        ("P0", "工程、九模块壳层、领域模型、合成目标", "x64 构建、导航和状态夹具；地址/偏移往返。"),
        ("P1", "通用 Win32 扫描、地址表、冻结", "扫描→结果→地址表→写入/全停闭环。"),
        ("P2", "统一十六进制、文件源、进程源", "本地文件与进程内存均可打开、查找、编辑、撤销、提交。"),
        ("P3", "文件浏览/内容扫描、旧格式导入", "递归取消、长路径、大文件、原子保存、导入预览。"),
        ("P4", "图片、宏、设置", "图片管理/转换；安全宏；旧设置映射。"),
        ("P5", "适配器平台、速度、首个模拟器", "契约测试、崩溃隔离、精确版本适配和健康检查。"),
    ], [1100, 3650, 4610], font_size=9)
    h(doc, "15.1 关键验收", 2)
    bullets(doc, [
        "打开一个大于 4GB 的稀疏测试文件，跳转到高偏移、修改 4 字节、撤销、重做、另存并校验哈希。",
        "文件扫描在权限拒绝、符号链接循环、文件删除和取消场景中不崩溃且进度可解释。",
        "同一十六进制页切换到进程源后，插入/删除自动禁用；覆盖写入执行进程身份复验和读回。",
        "九模块页面、动作栏、快捷键和禁用原因均可在设计时状态中重现。",
    ], bullet)

    h(doc, "16. 风险与待定项")
    base.add_table(doc, ["编号", "事项", "处理"], [
        ("TBD-FMT", ".fpe/.mis/.mi0 精确格式", "建立样本库与差分测试；未确认字段不写回。"),
        ("TBD-EDITOR", "原版 TabEditor 的全部数据源语义", "以用户确认的文件功能为基线，同时保留进程源。"),
        ("TBD-SPEED", "旧速度控制机制", "不复用未知钩子；由模拟器适配器能力替代。"),
        ("TBD-GPE", "采集触发条件与部分字段", "先实现图片目录/命名/转换，再用合法测试素材验证采集。"),
        ("TBD-EMU", "首个模拟器与版本", "通用平台完成后由项目所有者选择。"),
    ], [1500, 3200, 4660], font_size=9)
    base.add_callout(doc, "可执行结论", "P0-P4 不依赖首个模拟器即可开始；P5 的具体适配器需要精确版本、SHA-256、合法测试游戏和已知数值。", fill="EAF4EE", label_color=base.GREEN)
    finish(doc, ARCH_OUT)


def build_code():
    doc, bullet, dec = new_doc("FPE2001-Remake 代码实施规格包 v2.1")
    cover(doc, "IMPLEMENTATION SPECIFICATION PACKAGE", "代码实施规格包 v2.1",
          "覆盖九模块、文件/进程统一十六进制工作台与逐模拟器适配",
          [("技术栈", ".NET 10 LTS、C# 14、WPF、x64、SQLite、命名管道"),
           ("工程范围", "Memory / Scan / AddressBook / BinaryEditor / FileWorkspace / Media / Automation / Speed / Adapters"),
           ("安全边界", "本地离线；无驱动、注入或反作弊绕过"),
           ("版本", "v2.1 · 2026-08-16")],
          "本规格可分票交给其他模型实施；文件编辑器不再是 P4 附件，而是 P2 核心能力。")
    impl.add_contents(doc, ["实施合同", "仓库与程序集", "公共领域模型", "统一字节源", "文件十六进制编辑器",
                            "进程内存编辑", "文件浏览与扫描", "进程扫描与地址表", "图片模块", "宏与热键", "速度能力",
                            "适配器与 IPC", "数据库与文件格式", "安全与诊断", "测试矩阵", "P0-P6 任务单", "模型提示词与待定项"], dec[0])

    h(doc, "1. 实施合同")
    base.add_callout(doc, "MUST", "所有九模块都有路由、ViewModel、应用服务和设计时状态；M03 同时支持 File 与 ProcessMemory；M04 能浏览、筛选、递归扫描、取消并把结果送入 M03。")
    bullets(doc, [
        "Release win-x64 构建无警告；核心测试通过；退出后无孤儿 Broker/AdapterHost/扫描任务。",
        "地址、文件偏移和领域长度用 ulong；OS 调用边界可检查转换为 nuint。",
        "文件保存必须支持外部变化检测、备份和原子替换；不允许直接原地写坏原文件。",
        "进程写入必须复验 ProcessIdentity、可选写前比较和读回校验。",
        "长任务全部支持 CancellationToken、阶段进度和重复点击保护。",
    ], bullet)

    h(doc, "2. 仓库与程序集")
    code(doc, """
src/
  Fpe2001Remake.UI
  Fpe2001Remake.Application
  Fpe2001Remake.Domain
  Fpe2001Remake.Contracts
  Fpe2001Remake.Memory.Win32
  Fpe2001Remake.Scan
  Fpe2001Remake.AddressBook
  Fpe2001Remake.BinaryEditor.Core
  Fpe2001Remake.ByteSources.File
  Fpe2001Remake.ByteSources.Process
  Fpe2001Remake.FileWorkspace
  Fpe2001Remake.MediaWorkspace
  Fpe2001Remake.Automation
  Fpe2001Remake.SpeedControl
  Fpe2001Remake.Storage
  Fpe2001Remake.LegacyImport
  Fpe2001Remake.AdapterSdk
  Fpe2001Remake.AdapterHost
  Fpe2001Remake.Broker
adapters/
  GenericWin32 / LegacyIni / Libretro / <Emulator>
tests/
  Unit / Contract / Integration / EndToEnd / SyntheticTarget.x86 / SyntheticTarget.x64
""")
    base.add_table(doc, ["程序集", "职责", "禁止"], [
        ("BinaryEditor.Core", "分页缓存、游标/选择、搜索、变更集、撤销、书签、编码", "直接 WPF、直接 Win32、直接文件对话框。"),
        ("ByteSources.File", "长路径、64 位偏移、共享模式、外部变化、保存策略", "解析 UI 状态。"),
        ("ByteSources.Process", "逻辑地址解析、区域读取、覆盖写入", "插入/删除、持久化宿主地址。"),
        ("FileWorkspace", "目录枚举、筛选、递归内容扫描、结果", "默认跟随 reparse point。"),
        ("Automation", "热键、动作序列、目标/焦点守卫、急停", "全局任意进程控制。"),
        ("SpeedControl", "能力协商与值域映射", "通用时间钩子。"),
    ], [2100, 4200, 3060], font_size=8.9)

    h(doc, "3. 公共领域模型")
    code(doc, """
public enum AddressSpace { HostVirtual, GuestPhysical, GuestVirtual, ModuleRelative, PointerChain }
public enum ByteSourceKind { File, ProcessMemory, Snapshot }

public readonly record struct BytePosition(ByteSourceKind SourceKind, ulong Offset);
public readonly record struct ByteRange(ulong Offset, ulong Length);
public readonly record struct LogicalAddress(
    AddressSpace Space, string DomainId, ulong Value,
    Endianness Endianness, byte PointerWidthBits);

public sealed record ByteSourceIdentity(
    Guid InstanceId, ByteSourceKind Kind, string DisplayName,
    ulong? Length, string VersionToken);
""")
    base.add_callout(doc, "VersionToken", "文件源使用路径规范化+长度+LastWriteUtc+FileId 的稳定标记；进程源使用 PID+StartTimeUtc+适配器会话。任何提交前必须复验。")

    h(doc, "4. 统一字节源")
    code(doc, """
public interface IByteSource : IAsyncDisposable
{
    ByteSourceIdentity Identity { get; }
    ByteSourceCapabilities Capabilities { get; }
    ValueTask<ByteReadResult> ReadAsync(ulong offset, Memory<byte> destination, CancellationToken ct);
}

public interface IEditableByteSource : IByteSource
{
    ValueTask<CommitResult> CommitAsync(ChangeSet changes, CommitOptions options, CancellationToken ct);
}

[Flags]
public enum ByteSourceCapabilities
{
    Read = 1, Overwrite = 2, Insert = 4, Delete = 8,
    KnownLength = 16, AtomicCommit = 32, Refresh = 64
}
""")
    bullets(doc, [
        "ReadAsync 允许部分读取并返回 CompletedLength、UnavailableRanges 和可重试错误。",
        "编辑器页面大小默认 64 KiB，缓存上限默认 64 MiB；按 LRU 淘汰未修改页。",
        "所有 UI 结果携带 generationId；来源刷新或切换后丢弃过期响应。",
        "文件源支持 Read/Overwrite/Insert/Delete；进程源只支持 Read/Overwrite/Refresh。",
    ], bullet)

    h(doc, "5. 文件十六进制编辑器")
    h(doc, "5.1 变更集与撤销", 2)
    code(doc, """
public abstract record ByteEdit(ulong Offset);
public sealed record OverwriteEdit(ulong Offset, byte[] Before, byte[] After) : ByteEdit(Offset);
public sealed record InsertEdit(ulong Offset, byte[] Data) : ByteEdit(Offset);
public sealed record DeleteEdit(ulong Offset, byte[] Deleted) : ByteEdit(Offset);

public sealed record ChangeSet(
    Guid Id, string BaseVersionToken,
    IReadOnlyList<ByteEdit> Edits, ulong ResultLength);
""")
    bullets(doc, [
        "编辑动作合并窗口默认 500ms；连续同类覆盖可合并为一个撤销单元。",
        "撤销栈默认内存配额 128MiB，超限后把旧变更块压缩落盘；用户必须看到可撤销范围。",
        "插入/删除使用区间映射，不为大文件复制完整缓冲；修改高亮由区间树查询。",
        "关闭有修改的标签页时提供保存、另存、放弃、取消；默认焦点为取消。",
    ], bullet)
    h(doc, "5.2 安全提交", 2)
    numbers(doc, [
        "复验 VersionToken；不一致则进入三方选择：重新加载、另存为、查看外部变化。",
        "计算目标目录可用空间；默认创建 .bak 或时间戳备份。",
        "在同目录创建临时文件，顺序流式应用变更；Flush(true)。",
        "校验长度和可选 SHA-256；使用 ReplaceFile/MoveFileEx 等价原子替换。",
        "重新打开并验证新 VersionToken；失败保留临时文件与备份，给出恢复路径。",
    ], dec[1])
    h(doc, "5.3 搜索与编码", 2)
    base.add_table(doc, ["模式", "输入", "规则"], [
        ("Hex", "DE AD ?? EF", "支持掩码；不跨不可读洞；块间保留 patternLength-1。"),
        ("Text", "文本+编码", "ASCII/UTF-8/UTF-16LE/GBK；大小写可选。"),
        ("Integer", "类型/值/端序", "Int/UInt 8-64；精确/范围。"),
        ("Float", "Float32/64", "epsilon；NaN 默认不匹配。"),
    ], [1500, 2500, 5360], font_size=9)

    h(doc, "6. 进程内存编辑")
    bullets(doc, [
        "来源可由扫描结果、地址表或手工逻辑地址打开；标题显示目标、域和 ProcessIdentity。",
        "不可读字节显示 ??；选择可跨可读区，但提交必须分段并逐段验证。",
        "插入/删除控件隐藏；覆盖编辑默认关闭，开启后整页显示持续警示状态。",
        "刷新创建新 generation；有未提交修改时刷新需确认或先导出补丁。",
        "写入使用 ExpectedCurrentValue 和 ReadBackExact；目标重启后所有提交拒绝。",
    ], bullet)

    h(doc, "7. 文件浏览与扫描")
    code(doc, """
public sealed record FileScanRequest(
    string RootPath, bool Recursive, bool FollowReparsePoints,
    FileFilter Filter, ContentPredicate Predicate,
    int MaxConcurrency, ulong? MaxBytes, int MaxMatchesPerFile);

public sealed record FileMatch(
    string Path, ulong Offset, int Length,
    string Preview, string VersionToken, string MatchKind);
""")
    bullets(doc, [
        "路径使用长路径兼容 API；默认 FollowReparsePoints=false；检测目录 FileId 防循环。",
        "目录枚举与文件内容读取分离；默认并发 min(4, CPU)，机械盘可配置为 1。",
        "文件变化、删除、共享冲突和 AccessDenied 计入 SkipReason，不让整个任务失败。",
        "结果点击以保存的 VersionToken 打开；版本变化时重新定位并警告，不盲跳旧偏移。",
        "停止按钮 1 秒内进入 Cancelled；关闭所有流和映射；不留下临时索引。",
    ], bullet)

    h(doc, "8. 进程扫描与地址表")
    base.add_table(doc, ["接口", "责任"], [
        ("IMemoryRegionProvider", "VirtualQueryEx 区域规划、保护过滤、checked 推进。"),
        ("IScanEngine", "首次/再次扫描、Mission、进度、取消、FPSN 快照。"),
        ("IAddressBookService", "条目、分组、注释、值刷新、旧 .fpe 导入/导出。"),
        ("IFreezeService", "优先队列、条件、间隔、失败熔断、全局停止。"),
    ], [2500, 6860], font_size=9.2)
    bullets(doc, ["沿用 v1 的 FPSN v1 快照格式与扫描比较器。", "地址表保存逻辑地址，不跨会话持久化宿主地址。", "从地址表“编辑”打开 M03 的 ProcessMemoryByteSource。"], bullet)

    h(doc, "9. 图片模块")
    code(doc, """
public interface IMediaWorkspace
{
    IAsyncEnumerable<MediaItem> EnumerateAsync(string folder, CancellationToken ct);
    Task<MediaItem> RenameAsync(MediaIdentity item, string newName, CancellationToken ct);
    Task ExportAsync(MediaIdentity item, MediaExportOptions options, CancellationToken ct);
    Task DeleteAsync(MediaIdentity item, DeletePolicy policy, CancellationToken ct);
}
""")
    bullets(doc, ["支持 BMP/JPG/GIF 读取与导出；解码在后台并限制像素数。", "连续编号与 x2 仅改变命名/预览策略，不隐式覆盖原图。", "删除优先回收站；导出显示格式、尺寸、JPG 质量和目标路径。"], bullet)

    h(doc, "10. 宏与热键")
    base.add_table(doc, ["对象", "字段"], [
        ("Macro", "Id、Name、TargetBinding、Hotkey、Steps、Repeat、Timeout、Enabled"),
        ("Step", "KeyDown/Up、MouseLeft/Right/Double、Delay、WaitForeground、Comment"),
        ("TargetBinding", "进程身份规则、窗口类/标题规则、适配器会话可选"),
    ], [2200, 7160], font_size=9.2)
    bullets(doc, [
        "优先 RegisterHotKey；只有录制/兼容场景才使用低级钩子并显示持续状态。",
        "运行前 3 秒倒计时可选；焦点不匹配、目标退出、超时或 Ctrl+Shift+F12 立即停止。",
        "宏日志默认只记录步骤类型和时间，不记录用户输入文本。",
    ], bullet)

    h(doc, "11. 速度能力")
    code(doc, """
public interface ISpeedControlCapability
{
    SpeedRange Range { get; }
    ValueTask<SpeedState> GetAsync(CancellationToken ct);
    ValueTask<SpeedState> SetEnabledAsync(bool enabled, CancellationToken ct);
    ValueTask<SpeedState> SetMultiplierAsync(double multiplier, CancellationToken ct);
}
""")
    bullets(doc, ["值域与步进来自适配器；UI 显示 0.25x/0.5x/1x/2x 等推荐刻度。", "没有能力时返回 UnsupportedCapability，不尝试通用 DLL 钩子。", "目标暂停、快进和速度倍率是独立能力，不互相推断。"], bullet)

    h(doc, "12. 适配器与 IPC")
    bullets(doc, [
        "沿用 v1 的 IEmulatorAdapter、命名管道帧、protocolVersion、messageId、correlationId、deadlineUtc 和 nonce。",
        "新增 capability.speed.get/set、capability.media.capture（可选）和 automation.target.validate。",
        "地址序列化为 0x+16 位字符串；文件偏移同样使用字符串。",
        "AdapterHost 崩溃只断开目标能力；文件编辑和未提交变更必须继续存在。",
    ], bullet)

    h(doc, "13. 数据库与文件格式")
    code(doc, """
CREATE TABLE editor_sessions(
  id TEXT PRIMARY KEY, source_kind INTEGER NOT NULL,
  source_identity_json TEXT NOT NULL, cursor_offset TEXT NOT NULL,
  selection_length TEXT NOT NULL, encoding TEXT NOT NULL,
  overwrite_mode INTEGER NOT NULL, updated_utc TEXT NOT NULL);

CREATE TABLE bookmarks(
  id TEXT PRIMARY KEY, editor_session_id TEXT NOT NULL,
  offset_hex TEXT NOT NULL, label TEXT NOT NULL, color_key TEXT,
  FOREIGN KEY(editor_session_id) REFERENCES editor_sessions(id) ON DELETE CASCADE);

CREATE TABLE file_scan_history(
  id TEXT PRIMARY KEY, request_json TEXT NOT NULL,
  state INTEGER NOT NULL, files_visited INTEGER NOT NULL,
  bytes_read TEXT NOT NULL, matches INTEGER NOT NULL, updated_utc TEXT NOT NULL);
""")
    bullets(doc, ["大匹配结果使用磁盘快照，不把每条 FileMatch 写成 SQLite 行。", "未保存文件变更可按用户设置恢复到加密本地临时区；默认关闭。", "旧格式导入只写新表，不原地修改旧文件。"], bullet)

    h(doc, "14. 安全与诊断")
    base.add_table(doc, ["风险", "控制"], [
        ("文件损坏", "默认只读、备份、同目录临时文件、原子替换、哈希、外部变化检测。"),
        ("目录越界", "用户选择根目录；路径规范化；默认不跟随 reparse point。"),
        ("误写进程", "最小权限、身份复验、写前比较、读回、全停。"),
        ("宏失控", "目标绑定、焦点守卫、超时、可见状态、急停。"),
        ("恶意媒体", "像素/帧/解码时间限制；异常隔离。"),
    ], [1900, 7460], font_size=9.2)

    h(doc, "15. 测试矩阵")
    base.add_table(doc, ["ID", "场景", "通过标准"], [
        ("HEX-001", ">4GB 稀疏文件高偏移", "跳转、读取、覆盖、撤销、另存后字节与长度正确。"),
        ("HEX-002", "插入/删除组合", "区间映射、游标、书签和结果长度正确。"),
        ("HEX-003", "外部修改", "提交被阻止；可重新加载或另存；原文件不损坏。"),
        ("HEX-004", "进程源", "插入/删除禁用；重启后写入被拒绝。"),
        ("FILE-001", "符号链接循环", "不循环；记录跳过原因。"),
        ("FILE-002", "取消大目录扫描", "1 秒内取消；句柄释放；结果保持一致。"),
        ("FILE-003", "匹配文件变化", "打开时警告并重新定位。"),
        ("MAC-001", "目标失焦/急停", "不再发送动作；500ms 内进入停止。"),
        ("SPD-001", "无速度能力", "页面说明不可用，不加载未知钩子。"),
        ("E2E-001", "原版核心流程", "扫描→地址表→编辑；文件→扫描→编辑→备份保存。"),
    ], [1300, 2700, 5360], font_size=8.9)

    h(doc, "16. P0-P6 任务单")
    base.add_table(doc, ["阶段", "任务包", "完成定义"], [
        ("P0", "九模块壳层、Domain、构建、合成目标/文件", "模块顺序、路由、64 位往返、CI。"),
        ("P1", "Memory/Scan/AddressBook/Freeze", "进程闭环与自动化验收。"),
        ("P2", "BinaryEditor.Core + File/Process sources", "HEX-001 至 HEX-004。"),
        ("P3", "FileWorkspace + LegacyImport", "FILE-001 至 FILE-003；旧配置预览。"),
        ("P4", "Media + Automation + Settings", "图片格式、宏急停、设置迁移。"),
        ("P5", "AdapterHost + IPC + Speed", "契约、崩溃隔离、无能力状态。"),
        ("P6", "首个模拟器适配", "精确版本、3 个已知地址、速度/暂停能力按需。"),
    ], [1000, 3600, 4760], font_size=9)

    h(doc, "17. 模型提示词与待定项")
    code(doc, """
只实现任务票 {TICKET_ID}，以 v2.1 规格的 {SECTIONS} 为验收基线。
先写失败测试，再实现最小变更；不得删除九模块入口；不得把 File 编辑降为 ProcessMemory 编辑的别名。
地址、文件偏移和长度不得降为 32 位。文件提交必须保留备份/原子性；进程提交必须复验身份。
运行受影响测试和 Release win-x64 构建，输出修改文件、测试结果、风险和 SPEC-GAP。
若缺失信息会改变公共契约，停止并标记 SPEC-GAP，不自行发明。
""")
    bullets(doc, ["TBD-FMT：旧二进制格式精确字段。", "TBD-GPE：采集触发语义。", "TBD-EMU：首个模拟器和版本。", "TBD-SIGN：发布签名与第三方适配器策略。"], bullet)
    finish(doc, CODE_OUT)


def build_uiux():
    doc, bullet, dec = new_doc("FPE2001-Remake UI/UX 实施规格 v2.1")
    cover(doc, "DESKTOP UI / UX IMPLEMENTATION SPECIFICATION", "UI/UX 实施规格 v2.1",
          "以原版九模块、动作栏和快捷键为骨架的 Windows 11 桌面设计",
          [("导航", "横向九模块带：扫描、地址表、十六进制、文件、图片、宏、速度、设置、关于"),
           ("核心变化", "M03 支持文件/进程双来源；M04 保留目录浏览与内容扫描"),
           ("视觉", "低饱和蓝灰、暖灰背景、等宽地址、标准/紧凑密度"),
           ("版本", "v2.1 · 2026-08-16")],
          "设计目标：老用户能按原顺序找到功能，新用户能清楚区分目标、数据源、地址/偏移和未保存修改。")
    impl.add_contents(doc, ["设计原则", "应用壳与原版导航", "公共动作栏与快捷键", "M01 扫描", "M02 地址表",
                            "M03 十六进制编辑", "M04 文件", "M05 图片", "M06 宏", "M07 速度", "M08 设置", "M09 关于",
                            "跨模块流程", "视觉令牌", "MVVM 映射", "状态、错误与无障碍", "性能与验收"], dec[0])

    h(doc, "1. 设计原则")
    bullets(doc, [
        "原版可识别性：九模块名称、顺序和主要 F 键动作保留；图标更新但语义不改。",
        "目标与来源可见：标题栏显示当前进程/模拟器；编辑器标签明确标注 File 或 Memory。",
        "修改可见：文件未保存、进程待写入、冻结运行、宏运行和速度非 1x 都有持续状态。",
        "高密度但不拥挤：表格紧凑、卡片少、主要动作在页面顶部，不用大型仪表盘。",
        "安全默认：文件只读打开、进程编辑关闭、宏停止、速度 1x、递归扫描关闭。",
    ], bullet)

    h(doc, "2. 应用壳与原版导航")
    base.add_table(doc, ["区域", "规格", "内容"], [
        ("标题栏", "46 DIP", "FPE2001-Remake、当前目标、连接/冻结/宏/速度状态、窗口按钮。"),
        ("模块带", "64-72 DIP", "九模块横向排列；图标+文字；保留原顺序；活动项低饱和填充。"),
        ("动作栏", "44 DIP", "当前模块动作；文字、图标和 F 键提示；危险动作分组。"),
        ("内容区", "自适应", "模块页面；数据表虚拟化；局部侧栏可折叠。"),
        ("状态栏", "32-36 DIP", "目标/来源、任务进度、匹配/候选、修改、冻结、宏、速度、急停。"),
    ], [1600, 1600, 6160], font_size=9.2)
    module_table(doc)
    bullets(doc, ["默认窗口 1440×900；最小 1100×720；200% DPI 无裁切。", "窄宽度下模块带水平滚动，不改变模块顺序；关于可进入更多，但设置保持可达。", "Ctrl+1…9 切换模块；Ctrl+Shift+F12 全局停止冻结、宏和速度临时覆盖。"], bullet)

    h(doc, "3. 公共动作栏与快捷键")
    base.add_table(doc, ["模块", "原快捷键映射", "v2.1 动作"], [
        ("扫描", "F2-F11、Enter、Esc", "清空、新任务、命名、Mission、计算、添加、列表、编辑、撤销、扫描、停止。"),
        ("地址表", "F1-F8", "添加、删除、清空、修改、编辑、写入、加载、保存。"),
        ("编辑器", "F1-F6", "打开、查找、撤销、书签、刷新、注释；新增保存/另存。"),
        ("文件", "F1-F6", "子目录、全部、刷新、扫描、编辑/保存、停止。"),
        ("图片", "F1-F4", "载入、命名、删除、编辑。"),
        ("宏", "F1-F6", "设置、删除、更新、鼠标左/右/双击。"),
    ], [1500, 1900, 5960], font_size=9)
    p(doc, "F 键提示显示在 tooltip 和命令面板中；若系统/辅助技术占用快捷键，命令仍可通过按钮执行。", color=base.MUTED)

    h(doc, "4. M01 扫描")
    base.add_table(doc, ["区域", "内容"], [
        ("任务栏", "Mission 列表、新建/命名、加载/保存、撤销；保留原任务概念。"),
        ("目标条件", "当前目标、模拟器/内存域、输入值、首次/再次扫描条件、8/16/32/64/Float/Bytes/Text。"),
        ("扫描控制", "扫描/停止、字节序、对齐、逻辑范围、未知值、进度。"),
        ("候选结果", "逻辑地址、当前/上次值、宿主证据、选中、添加到地址表、在编辑器打开。"),
        ("内存摘要", "计划/已读字节、区域、跳过、候选、快照、耗时。"),
    ], [1900, 7460], font_size=9.2)
    bullets(doc, ["首次完成后主按钮改为“再次扫描”，旁边保留“新任务”。", "Mission 与扫描历史分开：Mission 是用户命名工作，历史是不可变快照。", "取消不是错误；保留上个完整快照。"], bullet)

    h(doc, "5. M02 地址表")
    base.add_table(doc, ["列", "规则"], [
        ("启用", "复选/冻结状态；停止、等待条件、失败有文本。"),
        ("描述", "名称、分组、游戏、作者/制作者、注释。"),
        ("地址", "逻辑地址优先；宿主地址在详情；等宽字体。"),
        ("值", "当前值、写入值、类型、显示格式。"),
        ("锁定/条件", "Always/Equals/NotEquals/InRange/Changed、间隔、失败计数。"),
        ("操作", "修改、编辑、立即写入、回滚、删除。"),
    ], [1900, 7460], font_size=9.2)
    bullets(doc, ["顶部保留游戏、制作者、注释和表启用字段；加载/保存 .fpe 使用预览与兼容警告。", "双击行打开编辑抽屉；“编辑内存”跳转 M03 并创建 ProcessMemory 标签。"], bullet)

    h(doc, "6. M03 十六进制编辑")
    base.add_callout(doc, "两种来源", "顶部标签必须显示“文件”或“进程内存”。文件支持插入/删除与保存；进程只支持覆盖写入与刷新。")
    base.add_table(doc, ["区域", "文件来源", "进程来源"], [
        ("标签", "文件名、路径、* 未保存、版本标记", "目标名、域、起址、进程启动时间"),
        ("动作栏", "打开、保存、另存、查找、撤销/重做、书签、刷新、注释", "打开地址、查找、撤销待提交、书签、刷新、注释、提交写入"),
        ("网格", "64 位文件偏移、16 字节、文本、修改高亮", "逻辑/宿主地址、不可读 ??、待写入高亮"),
        ("编辑", "只读/覆盖/插入；删除；粘贴字节", "只读/覆盖；禁止插入删除"),
        ("状态栏", "偏移、选择长度、总长度、编码、模式、修改数、SHA", "地址、选择、区域保护、目标身份、待写字节"),
    ], [1500, 3900, 3960], font_size=8.8)
    h(doc, "6.1 查找面板", 2)
    bullets(doc, ["模式：Hex、Text、Integer、Float；编码、端序、范围、方向、大小写。", "结果显示偏移/地址、预览、类型；点击定位；大结果窗口化。", "替换仅文件来源可批量执行；进程来源逐项确认，不提供全局替换。"], bullet)
    h(doc, "6.2 保存与关闭", 2)
    bullets(doc, ["保存前显示差异摘要、外部修改、备份、输出长度和磁盘空间。", "关闭未保存标签：保存/另存/放弃/取消；默认取消。", "保存失败保留变更和恢复路径，不自动重试覆盖。"], bullet)

    h(doc, "7. M04 文件")
    base.add_table(doc, ["区域", "控件与行为"], [
        ("路径栏", "驱动器、面包屑路径、上级、刷新、系统选择器。"),
        ("目录树", "惰性加载；访问拒绝/联接/链接有标记。"),
        ("文件表", "名称、大小、类型、修改时间、扫描状态；扩展名/全部文件过滤。"),
        ("扫描抽屉", "Hex/Text/Integer/Float、递归、大小/时间/扩展名、上限、开始/停止。"),
        ("结果", "路径、64 位偏移、预览、匹配类型；打开到 M03。"),
        ("状态", "文件数、字节、匹配、跳过、错误、耗时、取消。"),
    ], [1900, 7460], font_size=9.2)
    bullets(doc, ["默认当前目录、非递归、不跟随链接、只读。", "“编辑/保存(F5)”在有选择文件时进入 M03；保存仍由编辑器负责。", "拖入文件等价于在 M03 只读打开；拖入目录等价于导航，不自动扫描。"], bullet)

    h(doc, "8. M05 图片")
    base.add_table(doc, ["区域", "内容"], [
        ("目录/文件", "图片目录、列表、缩略图、格式/尺寸/大小。"),
        ("预览", "适应、1:1、x2、缩放、透明背景。"),
        ("命名/采集", "ON/OFF、编号、名称、路径、++ 连续编号；能力未知时显示兼容说明。"),
        ("编辑/导出", "重命名、删除到回收站、BMP/JPG/GIF、JPG 质量。"),
    ], [2000, 7360], font_size=9.2)

    h(doc, "9. M06 宏")
    base.add_table(doc, ["区域", "内容"], [
        ("宏列表", "名称、热键、目标、步骤数、状态、最近运行。"),
        ("步骤表", "序号、动作、参数、延迟、目标条件、注释。"),
        ("动作栏", "设置/删除/更新、鼠标左/右/双击、键盘、延迟、等待前台。"),
        ("运行", "倒计时、当前步骤、剩余时间、停止；焦点/目标守卫。"),
    ], [2000, 7360], font_size=9.2)
    base.add_callout(doc, "安全", "宏运行时标题栏和状态栏持续显示；Ctrl+Shift+F12 停止；目标不匹配时不发送输入。", fill="FFF2F1", label_color=base.RED)

    h(doc, "10. M07 速度")
    bullets(doc, ["保留原版开关+滑杆的核心形式。", "显示当前倍率、推荐刻度、恢复 1x、适配器/目标、能力证据。", "无能力时控件禁用但页面保留，解释为什么不可用及如何选择适配器。", "离开页面不恢复；断开目标、急停或会话结束恢复默认状态。"], bullet)

    h(doc, "11. M08 设置")
    base.add_table(doc, ["分组", "原版映射与新增"], [
        ("扫描", "内存上限→自动/自定义 64 位范围；块大小、并发、对齐。"),
        ("文件与编辑", "临时目录、备份、撤销配额、缓存、默认编码、只读打开。"),
        ("写入与锁定", "锁定间隔、读回校验、全局停止快捷键。"),
        ("图片/SPE", "外部查看器、导出格式、质量、连续编号。"),
        ("宏", "热键、倒计时、目标守卫、低级钩子开发者开关。"),
        ("兼容/隐私", "旧配置预览、日志、诊断脱敏、开发者模式。"),
    ], [2100, 7260], font_size=9.2)

    h(doc, "12. M09 关于")
    bullets(doc, ["产品版本、协议版本、架构、更新渠道。", "原版来源与重构说明、第三方许可证、官方适配器列表。", "打开日志目录、复制系统信息、导出诊断、隐私说明。", "明确本地离线用途和不支持受保护/联网对战。"], bullet)

    h(doc, "13. 跨模块流程")
    base.add_table(doc, ["起点", "动作", "终点/上下文"], [
        ("扫描候选", "添加", "地址表；保留类型、域、端序、Proof。"),
        ("扫描候选/地址表", "编辑", "十六进制 Memory 标签；定位逻辑地址。"),
        ("文件列表/扫描结果", "编辑/打开", "十六进制 File 标签；定位 64 位偏移。"),
        ("图片文件", "编辑", "图片模块或受控外部 SPE；不进入十六进制，除非“以二进制打开”。"),
        ("宏/速度", "急停", "全部停止冻结+宏+速度覆盖；目标仍可保持连接。"),
    ], [1800, 1900, 5660], font_size=9)

    h(doc, "14. 视觉令牌")
    base.add_table(doc, ["令牌", "值", "用途"], [
        ("Background", "#F3F4F3", "应用内容背景"),
        ("Surface", "#FBFBFA", "动作栏、表格、侧栏"),
        ("Chrome", "#E7EAEB", "标题栏、模块带"),
        ("Primary", "#566F75", "活动模块、主按钮、选择"),
        ("Text", "#273336 / #6E7A7D", "主/次文本"),
        ("Success/Warning/Danger", "#657B6D / #8A7757 / #875F5F", "低饱和状态；必须配文字"),
        ("UI font", "Segoe UI Variable / Microsoft YaHei UI", "12-14 DIP"),
        ("Mono", "Cascadia Mono / Consolas", "地址、偏移、字节、哈希"),
    ], [2300, 3000, 4060], font_size=9.1)

    h(doc, "15. MVVM 映射")
    base.add_table(doc, ["模块", "ViewModel", "关键服务"], [
        ("M01", "ScanWorkspaceViewModel", "IScanApplicationService, IMissionService"),
        ("M02", "AddressTableViewModel", "IAddressBookService, IFreezeService"),
        ("M03", "HexEditorWorkspaceViewModel", "IByteSourceFactory, IHexEditorService"),
        ("M04", "FileWorkspaceViewModel", "IFileCatalog, IFileScanService"),
        ("M05", "MediaWorkspaceViewModel", "IMediaWorkspace"),
        ("M06", "MacroWorkspaceViewModel", "IAutomationService, IHotkeyService"),
        ("M07", "SpeedViewModel", "ISpeedControlService"),
        ("M08", "SettingsViewModel", "ISettingsService, ILegacyImportService"),
        ("M09", "AboutViewModel", "IVersionService, IDiagnosticsService"),
    ], [1100, 3600, 4660], font_size=8.9)

    h(doc, "16. 状态、错误与无障碍")
    bullets(doc, [
        "每页至少具备 Loading、Empty、Ready、Running、Cancelled、Failed、Disconnected/Unavailable。",
        "编辑器另有 ReadOnly、Dirty、Saving、ExternalChanged、SaveConflict、Recovered。",
        "文件页另有 AccessDenied、ReparseSkipped、FileChanged、PartialResults。",
        "错误包含用户说明、恢复动作、技术详情和 operationId；不只显示 Win32 文本。",
        "模块带、动作栏、表格、网格、弹窗均可键盘操作；图标按钮有自动化名称。",
        "200% DPI、高对比度、减少动画和中文/英文资源均为验收项。",
    ], bullet)

    h(doc, "17. 性能与验收")
    p(doc, "设计交接：实现截图必须与 FPE2001-Remake UI Design Kit v2.1 对比；结构差异记录到 UI-DEVIATIONS.md，九模块功能入口不得省略。")
    base.add_table(doc, ["ID", "场景", "通过标准"], [
        ("UI-01", "九模块导航 / DPI / 键盘", "顺序不变；Ctrl+1…9；200% DPI 无裁切；全流程可不用鼠标。"),
        ("UI-02", "大文件", "高偏移可跳转；滚动流畅；UI 内存不随文件线性增长。"),
        ("UI-03", "文件修改", "* 标记、差异、关闭确认、外部变化冲突完整。"),
        ("UI-04", "目录扫描", "进度、跳过、取消、部分结果和打开编辑器完整。"),
        ("UI-05", "进程编辑", "File/Memory 来源不混淆；插入删除禁用；身份变化提示。"),
        ("UI-06", "宏/冻结/速度急停", "500ms 内视觉进入停止；无新动作。"),
    ], [1300, 2500, 5560], font_size=9)
    finish(doc, UI_OUT)


def main():
    build_architecture()
    build_code()
    build_uiux()
    print(ARCH_OUT)
    print(CODE_OUT)
    print(UI_OUT)


if __name__ == "__main__":
    main()
